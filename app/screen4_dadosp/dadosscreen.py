from kivymd.uix.screen import MDScreen
from kivymd.uix.floatlayout import MDFloatLayout
from kivymd.uix.label import MDLabel
from kivymd.uix.button import MDButton, MDButtonText, MDButtonIcon
from kivymd.uix.textfield import MDTextField, MDTextFieldHintText, MDTextFieldHelperText, MDTextFieldTrailingIcon
from kivymd.uix.filemanager import MDFileManager
from kivy.core.window import Window
from kivy.uix.boxlayout import BoxLayout
from app.screen4_dadosp.dados_function import extrair_dados_pdf
from app.screen3_insertpdf.pdf_function import *
import os
import re
from docx import Document

class DadosScreen(MDScreen):
    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.layout = MDFloatLayout()


        self.label = MDLabel(
            text="Dados do Proponente",
            halign="center",
            theme_text_color="Custom",
            text_color="yellow",
            size_hint=(1, 0.1),
            pos_hint={"center_x": 0.5, "y": 0.9},
        )
        self.layout.add_widget(self.label)

        self.proponente = MDTextField(
            MDTextFieldHintText(text="Proponente"),
            pos_hint={"center_x": 0.5, "y": 0.75},
            size_hint=(0.8, None),
            height=50,
        )
        self.layout.add_widget(self.proponente)

        self.cpf = MDTextField(
            MDTextFieldHintText(text="CPF"),
            pos_hint={"center_x": 0.5, "y": 0.65},
            size_hint=(0.8, None),
            height=50,
        )
        self.layout.add_widget(self.cpf)

        self.botao_selecionar = MDButton(
            pos_hint={"center_x": 0.5, "y": 0.5},
            on_release=self.abrir_seletor_pdf,
        )
        self.botao_selecionar.add_widget(MDButtonText(text="Selecionar PDF"))
        self.layout.add_widget(self.botao_selecionar)

        self.file_manager = MDFileManager(
            exit_manager=self.fechar_arquivo,
            select_path=self.on_pdf_selecionado,
            ext=[".pdf"]
        )        
        
        self.add_widget(self.layout)

    def abrir_seletor_pdf(self, *args):
        self.file_manager.show(os.getcwd())

    def fechar_arquivo(self, *args):
        self.file_manager.close()

    def on_pdf_selecionado(self, caminho_pdf):
        """
        Callback do FileManager. Recebe apenas 1 parâmetro,
        extrai nome/CPF do PDF selecionado e preenche os campos.
        """
        self.fechar_arquivo()
        nome, cpf = extrair_dados_pdf(caminho_pdf)
        self.proponente.text = nome
        self.cpf.text        = cpf

    def preencher_com_dados(self, caminho_car, caminho_cit):
        # chamada pelo go_next da tela anterior
        self.caminho_car = caminho_car
        self.caminho_cit = caminho_cit
        nome, cpf = extrair_dados_pdf(caminho_car)
        self.proponente.text = nome
        self.cpf.text = cpf

    def receber_dados(self):
        """
        Também um método da classe: gera o Word com os valores dos campos.
        """
        doc = getattr(self, 'doc', None)
        if not doc:
            modelo_path = os.path.join(os.getcwd(), "models", "MODELO_LAUDO.docx")
            doc = Document(modelo_path)
        nome = self.proponente.text
        cpf  = self.cpf.text

        modelo_path = os.path.join(os.getcwd(), "models", "MODELO_LAUDO.docx")
        doc = Document(modelo_path)

        # substituição em parágrafos
        for par in doc.paragraphs:
            for run in par.runs:
                run.text = run.text.replace("#PROPONENTE", nome)
                run.text = run.text.replace("#CPF_PROPONENTE", cpf)

        # substituição em tabelas
        for tabela in doc.tables:
            for linha in tabela.rows:
                for celula in linha.cells:
                    for p in celula.paragraphs:
                        for run in p.runs:
                            run.text = run.text.replace("#PROPONENTE", nome)
                            run.text = run.text.replace("#CPF_PROPONENTE", cpf)

        self.doc = doc
        nome_limpo = re.sub(r"[^\w\s-]", "", nome)
        saida = os.path.join(os.getcwd(), f"LAUDO DE AVALIAÇÃO {nome_limpo}.docx")
        doc.save(saida)
        print(f"✅ Documento salvo em: {saida}")
