import fitz  # PyMuPDF
import re
import os
from docx import Document

def extrair_dados_pdf(pdf_path):
    nome = ""
    cpf = ""

    try:
        doc = fitz.open(pdf_path)
        for page in doc:
            texto = page.get_text()

            match_nome = re.search(r"\bNome:\s*(.+)", texto)
            match_cpf = re.search(r"\bCPF:\s*(\d{3}\.\d{3}\.\d{3}-\d{2})", texto)

            if match_nome:
                nome = match_nome.group(1).strip()
            if match_cpf:
                cpf = match_cpf.group(1).strip()

            if nome and cpf:
                break

        doc.close()

        # Abre o modelo Word
        modelo_path = os.path.join(os.getcwd(), "models", "MODELO_LAUDO.docx")
        if not os.path.exists(modelo_path):
            raise FileNotFoundError(f"Modelo não encontrado: {modelo_path}")

        doc_word = Document(modelo_path)

        # Substitui nos parágrafos mantendo formatação
        for par in doc_word.paragraphs:
            for run in par.runs:
                if "#PROPONENTE" in run.text:
                    run.text = run.text.replace("#PROPONENTE", nome)
                if "#CPF_PROPONENTE" in run.text:
                    run.text = run.text.replace("#CPF_PROPONENTE", cpf)

        # Também verifica dentro de tabelas
        for tabela in doc_word.tables:
            for linha in tabela.rows:
                for celula in linha.cells:
                    for par in celula.paragraphs:
                        for run in par.runs:
                            if "#PROPONENTE" in run.text:
                                run.text = run.text.replace("#PROPONENTE", nome)
                            if "#CPF_PROPONENTE" in run.text:
                                run.text = run.text.replace("#CPF_PROPONENTE", cpf)

        # Salva com nome limpo
        nome_limpo = re.sub(r"[^\w\s-]", "", nome).replace(" ", "_")
        saida_path = os.path.join(os.getcwd(), f"LAUDO_{nome_limpo}.docx")
        doc_word.save(saida_path)
        print(f"✅ Documento salvo em: {saida_path}")

    except Exception as e:
        print(f"❌ Erro ao processar PDF: {e}")

    return nome or "Nome não encontrado", cpf or "CPF não encontrado"
