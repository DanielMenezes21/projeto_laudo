from docx import Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P

def extrair_textos_e_tabelas(doc_path):
    doc = Document(doc_path)
    body = doc.element.body
    elementos = list(body.iterchildren())

    resultados = {}
    i = 0
    while i < len(elementos):
        el = elementos[i]
        if isinstance(el, CT_P):
            par = next(p for p in doc.paragraphs if p._p == el)
            for run in par.runs:
                if run.font.color and run.font.color.rgb and str(run.font.color.rgb) == "FF0000":
                    texto_vermelho = par.text.strip()
                    tabela = None
                    if i + 1 < len(elementos) and isinstance(elementos[i+1], CT_Tbl):
                        tbl = next(t for t in doc.tables if t._tbl == elementos[i+1])
                        tabela = tbl
                    resultados[texto_vermelho] = tabela
        i += 1
    if not resultados:
        resultados["Nenhum texto em vermelho encontrado."] = None
    return resultados