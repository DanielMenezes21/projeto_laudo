from kivymd.uix.screen import MDScreen
from kivymd.uix.textfield import MDTextField, MDTextFieldHintText
from kivymd.uix.button import MDButton, MDButtonText, MDIconButton
from kivymd.uix.filemanager import MDFileManager
from kivymd.uix.boxlayout import MDBoxLayout
from kivymd.uix.scrollview import MDScrollView
from kivy.uix.widget import Widget
from datetime import datetime
import openpyxl
import re
from modules.data_folder import formatar_data
import os
from kivy.metrics import dp
from app.screen3_dadosp.screen3_1_matriculas.screen3_1_1_detalhes.matricula_detalhe_screen import MatriculaDetalheScreen
from app.screen3_dadosp.screen3_1_matriculas.matricula_function import go_back

class MatriculaScreen(MDScreen):
    def __init__(self, lista_dados_matriculas=None, **kwargs):
        super().__init__(**kwargs)
        self.lista_dados_matriculas = lista_dados_matriculas if lista_dados_matriculas is not None else []

        self.layout = MDBoxLayout(orientation="vertical", padding=20, spacing=20)
        self.scroll = MDScrollView()
        self.botoes_matriculas = MDBoxLayout(orientation="vertical", spacing=10, size_hint_y=None)
        self.botoes_matriculas.bind(minimum_height=self.botoes_matriculas.setter("height"))
        self.scroll.add_widget(self.botoes_matriculas)

        self.button_back = MDIconButton(
            icon="arrow-left",
            size_hint=(1, None),
            pos_hint={"x": 0.06, "y": 0.4},
            size=(dp(56), dp(56)),
            theme_text_color="Custom",
            text_color=(1, 1, 1, 1),
            on_release=lambda x: go_back(self),
        )

        self.botao_salvar = MDButton(
            MDButtonText(text="Salvar informações"),
            pos_hint={"center_x":0.9, "center_y":0.9},
            on_release=lambda x: self.salvar_dados()
        )

        self.layout.add_widget(self.button_back)
        self.layout.add_widget(self.scroll)
        self.layout.add_widget(self.botao_salvar)
        self.add_widget(self.layout)

        self.file_manager_excel = MDFileManager(
            exit_manager=self.fechar_filemanager,
            select_path=self.selecionar_arquivo,
            ext=[".xlsx", ".xls"],
            preview=False
        )

        self.file_manager_imagem = MDFileManager(
            exit_manager=self.fechar_filemanager,
            select_path=self.selecionar_arquivo,
            ext=[".png", ".jpg", ".jpeg"],
            preview=True  
        )

        self.matriculas = {}
        self.matricula_selecionada = ""
        self.tipo_planilha = ""

    def extrair_numero_matricula_excel(self, caminho_arquivo):
        wb = openpyxl.load_workbook(caminho_arquivo, data_only=True)
        aba = "AREA UTIL"
        if aba not in wb.sheetnames:
            print(f"{aba} não encontrado")
            return ""
        ws = wb[aba]
        pattern = re.compile(r"matr[ií]cula[\s\:\-\t]*([\d\.\,]+)", re.IGNORECASE)
        for row in ws.iter_rows(values_only=True):
            for idx, cell in enumerate(row):
                if isinstance(cell, str):
                    cell_limpa = cell.strip().replace('\n', '').replace('\r', '').replace('\t', ' ')
                    match = pattern.search(cell_limpa)
                    if match:
                        numero = match.group(1)
                        return numero
                    if "matr" in cell_limpa.lower() and idx + 1 < len(row):
                        prox = row[idx + 1]
                        if isinstance(prox, (int, float)):
                            numero = str(prox)
                            return numero
                if isinstance(cell, (int, float)) and idx > 0:
                    ant = row[idx - 1]
                    if isinstance(ant, str) and "matr" in ant.lower():
                        numero = str(cell)
                        return numero
        print("Número de matrícula não encontrado")
        return ""
    
    
    def receber_lat_long_pdf(self, latitude, longitude):
        """
        Preenche os campos de latitude e longitude das matrículas.
        """
        for i, nome_matricula in enumerate(self.matriculas):
            campos = self.matriculas[nome_matricula]["campos"]
            if isinstance(latitude, list) and i < len(latitude):
                campos["latitude"].text = str(latitude[i])
            else:
                campos["latitude"].text = str(latitude)
            if isinstance(longitude, list) and i < len(longitude):
                campos["longitude"].text = str(longitude[i])
            else:
                campos["longitude"].text = str(longitude)

    def extrair_valor_total_excel(self, caminho_arquivo):
        wb = openpyxl.load_workbook(caminho_arquivo, data_only=True)
        aba = "SANEAMENTO"
        if aba not in wb.sheetnames:
            print(f"{aba} não encontrado")
            return ""
        ws = wb[aba]
        for row in ws.iter_rows(values_only=True):
            for idx, cell in enumerate(row):
                if isinstance(cell, str) and cell.strip().lower() == "valor total":
                    for prox in row[idx+1:]:
                        if prox not in (None, "", "-"):
                            try:
                                valor = float(prox)
                                valor_formatado = f"R$ {valor:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
                            except Exception:
                                valor_formatado = str(prox)
                            print(f"Valor Total encontrado: {valor_formatado}")
                            return valor_formatado
        print("Valor Total não encontrado")
        return ""
    
    def extrair_valor_liq_excel(self, caminho_arquivo):
        aba = "LIQUIDAÇÃO"
        try:
            wb = openpyxl.load_workbook(caminho_arquivo, data_only=True)
        except Exception as e:
            print(f"Erro ao abrir a planilha: {e}")
            return ""

        if aba not in wb.sheetnames:
            print(f"Aba '{aba}' não encontrada")
            return ""

        ws = wb[aba]

        for row in ws.iter_rows(values_only=True):
            for idx, cell in enumerate(row):
                if isinstance(cell, str) and cell.strip().lower() == "valor de liquidação forçada":
                    for prox in row[idx+1:]:
                        if prox not in (None, "", "-"):
                            try:
                                valor = float(prox)
                                valor_formatado = f"R$ {valor:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
                            except Exception:
                                valor_formatado = str(prox)
                            print(f"Valor de Liquidação Forçada encontrado: {valor_formatado}")
                            return valor_formatado

        print("Valor de Liquidação Forçada não encontrado")
        return ""

    def extrair_area_total_excel(self, caminho_arquivo):
        aba = "AREA UTIL"
        try:
            wb = openpyxl.load_workbook(caminho_arquivo, data_only=True)
        except Exception as e:
            print(f"Erro ao abrir a planilha: {e}")
            return ""

        if aba not in wb.sheetnames:
            print(f"Aba '{aba}' não encontrada. Abas disponíveis: {wb.sheetnames}")
            return ""

        ws = wb[aba]

        for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
            for col_idx, cell in enumerate(row):
                if isinstance(cell, str):
                    texto = cell.strip().lower()
                    if texto == "area total":
                        try:
                            valor = ws.cell(row=row_idx + 2, column=col_idx + 1).value
                            if valor not in (None, "", "-"):
                                try:
                                    valor = float(valor)
                                    valor_formatado = f"{valor:.4f}".replace(".", ",")
                                except Exception:
                                    valor_formatado = str(valor)
                                print(f"✅ Área TOTAL encontrada: {valor_formatado}")
                                return valor_formatado
                            else:
                                print(f"⚠️ Célula abaixo de 'AREA TOTAL' está vazia.")
                        except Exception as e:
                            print(f"❌ Erro ao acessar célula abaixo: {e}")
                            return ""

        print("❌ 'AREA TOTAL' não encontrada na aba.")
        return ""
    
    def extrair_area_const_excel(self, caminho_arquivo):
        aba = "AREA UTIL"
        try:
            wb = openpyxl.load_workbook(caminho_arquivo, data_only=True)
        except Exception as e:
            print(f"Erro ao abrir a planilha: {e}")
            return ""

        if aba not in wb.sheetnames:
            print(f"Aba '{aba}' não encontrada. Abas disponíveis: {wb.sheetnames}")
            return ""

        ws = wb[aba]

        for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
            for col_idx, cell in enumerate(row):
                if isinstance(cell, str):
                    texto = cell.strip().lower()
                    if texto == "ÁREA CONSOLIDADA":
                        try:
                            valor = ws.cell(row=row_idx + 2, column=col_idx + 1).value
                            if valor not in (None, "", "-"):
                                try:
                                    valor = float(valor)
                                    valor_formatado = f"{valor:.4f}".replace(".", ",")
                                except Exception:
                                    valor_formatado = str(valor)
                                print(f"✅ ÁREA CONSOLIDADA encontrada: {valor_formatado}")
                                return valor_formatado
                            else:
                                print(f"⚠️ Célula abaixo de 'ÁREA CONSOLIDADA' está vazia.")
                        except Exception as e:
                            print(f"❌ Erro ao acessar célula abaixo: {e}")
                            return ""

        print("❌ 'ÁREA CONSOLIDADA' não encontrada na aba.")
        return ""

    def abrir_filemanager(self, matricula_nome, tipo):
        data = formatar_data()
        data_nome = datetime.now()
        mes = f'{data_nome.month:02d}. {data.split("de")[1].strip()}'
        self.matricula_selecionada = matricula_nome
        self.tipo_planilha = tipo

        initial_path = r"H:\\1. AVALIAÇÕES\\01. AVALIAÇÕES SICREDI\\01. RURAL"
        initial_path = os.path.join(initial_path, mes)
        if not os.path.exists(initial_path):
            initial_path = os.path.join(os.path.expanduser("~"), "Documents")
        self.root_path = initial_path
        self.current_path = self.root_path

        if tipo == "imagem":
            self.file_manager_imagem.show(initial_path)
        else:
            self.file_manager_excel.show(initial_path)

    def coletar_dados_matriculas(self):
        dados = []
        for nome_matricula, estrutura in self.matriculas.items():
            campos = estrutura["campos"]
            arquivos = estrutura["arquivos"]
            dados.append({
                "nome_imovel": campos["nome_imovel"].text,
                "matricula": campos["numero"].text,
                "valor_total": campos["valor"].text,
                "valor_liq": campos["valor_liq"].text,
                "latitude": campos["latitude"].text,
                "longitude": campos["longitude"].text,
                "proprietario": campos["proprietario"].text,
                "area_total": campos["area_total"].text,
                "area_consolidada": campos["area_consolidada"].text,
                "imagem": arquivos.get("imagem", ""),
                "planilha": arquivos.get("planilha", "")
            })
        print(f"Matrícula: {nome_matricula} | Proprietário coletado: {campos['proprietario'].text}")
        return dados

    def fechar_filemanager(self, *args):
        if getattr(self.file_manager_excel, "_window_manager", None):
            self.file_manager_excel.close()
        if getattr(self.file_manager_imagem, "_window_manager", None):
            self.file_manager_imagem.close()

    def selecionar_arquivo(self, caminho):
        self.fechar_filemanager()
        m = self.matricula_selecionada
        t = self.tipo_planilha
        self.matriculas[m]["arquivos"][t] = caminho
        print(f"📂 {t.upper()} associado a {m}: {caminho}")

        if t == "planilha":
            if caminho.endswith(".xlsx") or caminho.endswith(".xls"):
                try:
                    numero = self.extrair_numero_matricula_excel(caminho)
                    if numero:
                        self.matriculas[m]["campos"]["numero"].text = numero
                    valor_total = self.extrair_valor_total_excel(caminho)
                    if valor_total:
                        self.matriculas[m]["campos"]["valor"].text = str(valor_total)
                    valor_liq = self.extrair_valor_liq_excel(caminho)
                    if valor_liq:
                        self.matriculas[m]["campos"]["valor_liq"].text = str(valor_liq)
                    area_total = self.extrair_area_total_excel(caminho)
                    if area_total:
                        self.matriculas[m]["campos"]["area_total"].text = str(area_total)
                    area_consolidada = self.extrair_area_const_excel(caminho)
                    if area_consolidada:
                        self.matriculas[m]["campos"]["area_consolidada"].text = str(area_consolidada)
                except Exception as e:
                    print(f"❌ Erro ao processar planilha: {e}")
            else:
                print("❌ Arquivo selecionado não é uma planilha válida.")
        elif t == "imagem":
            print("🖼️ Imagem associada com sucesso.")

    def abrir_tela_detalhe(self, nome_matricula):
        nome_tela = f"detalhe_{nome_matricula.replace(' ', '_').lower()}"
        lista_dados = self.lista_dados_matriculas
        indice = None
        for i, d in enumerate(lista_dados):
            if d["nome_imovel"] == nome_matricula or d["matricula"] == nome_matricula:
                indice = i
                break
        if indice is None:
            indice = 0 

        if not self.manager.has_screen(nome_tela):
            nova_tela = MatriculaDetalheScreen(
                nome_matricula,
                lista_dados_matriculas=self.lista_dados_matriculas, 
                indice_matricula_atual=indice,
                name=nome_tela
            )
            self.manager.add_widget(nova_tela)
        self.manager.current = nome_tela

    def salvar_dados(self):
        for i, nome_matricula in enumerate(self.matriculas):
            campos = self.matriculas[nome_matricula]["campos"]
            arquivos = self.matriculas[nome_matricula]["arquivos"]
            if i < len(self.lista_dados_matriculas):
                self.lista_dados_matriculas[i].update({
                    "nome_imovel": campos["nome_imovel"].text,
                    "matricula": campos["numero"].text,
                    "valor_total": campos["valor"].text,
                    "valor_liq": campos["valor_liq"].text,
                    "latitude": campos["latitude"].text,
                    "area_total": campos["area_total"].text,
                    "area_consolidada": campos["area_consolidada"].text,
                    "longitude": campos["longitude"].text,
                    "proprietario": campos["proprietario"].text,
                    "imagem": self.matriculas[nome_matricula]["arquivos"].get("imagem", ""),
                    "planilha": arquivos.get("planilha", "")
                })
            else:
                self.lista_dados_matriculas.append({
                    "nome_imovel": campos["nome_imovel"].text,
                    "matricula": campos["numero"].text,
                    "valor_total": campos["valor"].text,
                    "valor_liq": campos["valor_liq"].text,
                    "latitude": campos["latitude"].text,
                    "longitude": campos["longitude"].text,
                    "area_total": campos["area_total"].text,
                    "area_consolidada": campos["area_consolidada"].text,
                    "proprietario": campos["proprietario"].text,
                    "imagem": self.matriculas[nome_matricula]["arquivos"].get("imagem", ""),
                    "planilha": arquivos.get("planilha", "")
                })
        print(f"🔄 Salvando dados para tela PDF: {self.lista_dados_matriculas}")
        tela_pdf = self.manager.get_screen('pdf')
        if hasattr(tela_pdf, "receber_dados_matriculas"):
            tela_pdf.receber_dados_matriculas(self.lista_dados_matriculas)


    def receber_dados_imoveis(self, imoveis, latitudes, longitudes, dados_completos=None, nomes_proprietarios=None):
        """Recebe os dados de imóveis de múltiplos PDFs
        
        Args:
            imoveis: Lista de nomes dos imóveis
            latitudes: Lista de latitudes
            longitudes: Lista de longitudes
            dados_completos: Dicionário completo com todos os dados (opcional)
            nomes_proprietarios: Lista de nomes dos proprietários (opcional)
        """
        self.dados_imoveis = dados_completos if dados_completos else {
            "imoveis": imoveis,
            "latitudes": latitudes,
            "longitudes": longitudes,
            "nomes_proprietarios": nomes_proprietarios or []
        }
        
        print(f"📌 Total de imóveis recebidos: {len(imoveis)}")
        if nomes_proprietarios:
            print(f"👤 Proprietários recebidos: {nomes_proprietarios}")

        if hasattr(self, 'matriculas'):
            for i, nome_matricula in enumerate(self.matriculas):
                campos = self.matriculas[nome_matricula]["campos"]
                if i < len(imoveis):
                    campos["nome_imovel"].text = imoveis[i]
                if i < len(latitudes):
                    campos["latitude"].text = str(latitudes[i])
                if i < len(longitudes):
                    campos["longitude"].text = str(longitudes[i])
                if nomes_proprietarios and i < len(nomes_proprietarios):
                    print(f"Preenchendo proprietário da matrícula {nome_matricula}: {nomes_proprietarios[i]}")
                    campos["proprietario"].text = nomes_proprietarios[i]

    def criar_botoes_para_matriculas(self, quantidade):
        self.botoes_matriculas.clear_widgets()
        self.matriculas.clear()

        main_layout = MDBoxLayout(orientation="vertical", spacing=25, size_hint_y=None)
        main_layout.bind(minimum_height=main_layout.setter('height'))

        for i in range(quantidade):
            self._adicionar_grupo_matricula(i+1, main_layout)

        add_layout = MDBoxLayout(orientation="horizontal", spacing=10, size_hint_y=None, height=dp(60))
        
        self.campo_adicional = MDTextField(
            MDTextFieldHintText(text="Qtd. adicional"),
            input_filter="int",
            size_hint_x=0.6,
        )
        
        btn_adicionar = MDButton(
            MDButtonText(text="Adicionar"),
            size_hint_x=0.4,
            on_release=self._adicionar_matriculas_extra
        )
        
        add_layout.add_widget(self.campo_adicional)
        add_layout.add_widget(btn_adicionar)

        self.botoes_matriculas.add_widget(main_layout)
        self.botoes_matriculas.add_widget(add_layout)

        self.lista_dados_matriculas.clear()
        for i in range(quantidade):
            self.lista_dados_matriculas.append({
                "nome_imovel": "",
                "matricula": "",
                "valor_total": "",
                "valor_liq": "",
                "latitude": "",
                "area_total": "",
                "area_consolidada": "",
                "longitude": "",
                "proprietario": "",
                "imagem": ""
            })

    def _adicionar_grupo_matricula(self, numero, layout_pai):
        """Cria todos os campos para uma matrícula, incluindo dropdown de múltipla escolha para proprietários"""
        from kivymd.uix.menu import MDDropdownMenu
        from kivymd.uix.boxlayout import MDBoxLayout
        from kivymd.uix.button import MDIconButton

        grupo = MDBoxLayout(orientation="vertical", spacing=15, size_hint_y=None, height=dp(920))

        dados_imovel = {}
        if hasattr(self, 'dados_imoveis'):
            if isinstance(self.dados_imoveis, list) and len(self.dados_imoveis) >= numero:
                dados_imovel = self.dados_imoveis[numero-1]
            elif isinstance(self.dados_imoveis, dict):
                dados_imovel = {
                    'nome_imovel': self.dados_imoveis.get('imoveis', [''])[min(numero-1, len(self.dados_imoveis.get('imoveis', [])))],
                    'latitude': self.dados_imoveis.get('latitudes', [''])[min(numero-1, len(self.dados_imoveis.get('latitudes', [])))],
                    'longitude': self.dados_imoveis.get('longitudes', [''])[min(numero-1, len(self.dados_imoveis.get('longitudes', [])))],
                    'proprietario': self.dados_imoveis.get('nomes_proprietarios', [''])[min(numero-1, len(self.dados_imoveis.get('nomes_proprietarios',[])))]
                }

        campo_nome = MDTextField(
            MDTextFieldHintText(text=f"Nome do Imóvel {numero}"),
            text=dados_imovel.get('nome_imovel', ''),
            size_hint_x=0.9
        )

        proponentes_dict = self.manager.get_screen('dados').proponentes
        nomes_proponentes = list(proponentes_dict.keys())

        linha_prop = MDBoxLayout(orientation="horizontal", spacing=10, size_hint_x=0.9)

        campo_prop = MDTextField(
            MDTextFieldHintText(text=f"Proprietário(s) do imóvel {numero}"),
            text=dados_imovel.get('proprietario', ''),
            size_hint_x=0.85,
            readonly=True
        )

        selecionados = []

        def atualizar_campo():
            campo_prop.text = ", ".join(selecionados)

        def toggle_proprietario(nome):
            if nome in selecionados:
                selecionados.remove(nome)
            else:
                selecionados.append(nome)
            atualizar_campo()

        menu_items = [
            {
                "text": nome,
                "on_release": lambda x=nome: toggle_proprietario(x)
            }
            for nome in nomes_proponentes
        ]

        btn_menu = MDIconButton(
            icon="menu-down",
            pos_hint={"center_y": 0.5},
            on_release=lambda x: menu.open()
        )

        menu = MDDropdownMenu(
            caller=btn_menu,
            items=menu_items,
            width_mult=4,
        )

        linha_prop.add_widget(campo_prop)
        linha_prop.add_widget(btn_menu)

        campo_matricula = MDTextField(
            MDTextFieldHintText(text=f"Nº da Matrícula {numero}"),
            size_hint_x=0.9
        )

        campo_valor_total = MDTextField(
            MDTextFieldHintText(text=f"Valor Total {numero}"),
            size_hint_x=0.9
        )

        campo_valor_liq = MDTextField(
            MDTextFieldHintText(text=f"Valor Líquido {numero}"),
            size_hint_x=0.9
        )

        campos_areas = MDBoxLayout(orientation="horizontal", spacing=15, size_hint_y=None, height=dp(180))

        campo_area_total = MDTextField(
            MDTextFieldHintText(text=f"Área total da matrícula {numero}"),
            size_hint_x=0.45
        )

        campo_area_cons = MDTextField(
            MDTextFieldHintText(text=f"Área Consolidada da matrícula {numero}"),
            size_hint_x=0.45
        )

        campos_areas.add_widget(campo_area_total)
        campos_areas.add_widget(campo_area_cons)

        coords = MDBoxLayout(orientation="vertical", spacing=15, size_hint_y=None, height=dp(180))
        campo_latitude = MDTextField(
            MDTextFieldHintText(text=f"Latitude {numero}"),
            text=str(dados_imovel.get('latitude', '')),
            size_hint_x=0.9
        )
        campo_longitude = MDTextField(
            MDTextFieldHintText(text=f"Longitude {numero}"),
            text=str(dados_imovel.get('longitude', '')),
            size_hint_x=0.9
        )
        coords.add_widget(campo_latitude)
        coords.add_widget(campo_longitude)

        btn_planilha = MDButton(
            MDButtonText(text="Selecionar Planilha"),
            on_release=lambda x, n=numero: self.abrir_filemanager(f"Matrícula {n}", "planilha")
        )

        btn_imagem = MDButton(
            MDButtonText(text="Selecionar Imagem"),
            on_release=lambda x, n=numero: self.abrir_filemanager(f"Matrícula {n}", "imagem")
        )

        btn_detalhes = MDButton(
            MDButtonText(text=f"Detalhes Matrícula {numero}"),
            on_release=lambda x, n=numero: self.abrir_tela_detalhe(f"Matrícula {n}")
        )

        grupo.add_widget(campo_nome)
        grupo.add_widget(linha_prop)  
        grupo.add_widget(campo_matricula)
        grupo.add_widget(campo_valor_total)
        grupo.add_widget(campo_valor_liq)
        grupo.add_widget(campos_areas)
        grupo.add_widget(coords)
        grupo.add_widget(btn_planilha)
        grupo.add_widget(btn_imagem)
        grupo.add_widget(btn_detalhes)

        self.matriculas[f"Matrícula {numero}"] = {
            "campos": {
                "proprietario": campo_prop,  
                "nome_imovel": campo_nome,
                "numero": campo_matricula,
                "valor": campo_valor_total,
                "valor_liq": campo_valor_liq,
                "area_total":campo_area_total,
                "area_consolidada": campo_area_cons,
                "latitude": campo_latitude,
                "longitude": campo_longitude
            },
            "arquivos": {}
        }

        layout_pai.add_widget(grupo)

    def _adicionar_matriculas_extra(self, *args):
        """Adiciona N matrículas extras conforme usuário solicitou"""
        if self.campo_adicional.text.isdigit():
            qtd = int(self.campo_adicional.text)
            start_num = len(self.matriculas) + 1
            
            for i in range(qtd):
                self._adicionar_grupo_matricula(start_num + i, self.botoes_matriculas.children[1])
            self.campo_adicional.text = "" 