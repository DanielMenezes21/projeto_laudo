import openpyxl
import re
import unicodedata

def go_back(self):
    self.manager.current_screen.manager.current = "dados"

def extrair_numero_matricula_excel(caminho_arquivo):
    wb = openpyxl.load_workbook(caminho_arquivo, data_only=True)
    aba = "AREA UTIL"
    if aba not in wb.sheetnames:
        print(f"{aba} não encontrado")
        return ""
    ws = wb[aba]
    pattern = re.compile(r"matr[ií]cula[\s\:\-\t]*([\d\.\,]+)", re.IGNORECASE)
    for row in ws.iter_rows(values_only=True):
        for idx, cell in enumerate(row):
            if isinstance(cell, str):
                cell_limpa = cell.strip().replace('\n', '').replace('\r', '').replace('\t', ' ')
                match = pattern.search(cell_limpa)
                if match:
                    numero = match.group(1)
                    return numero
                if "matr" in cell_limpa.lower() and idx + 1 < len(row):
                    prox = row[idx + 1]
                    if isinstance(prox, (int, float)):
                        numero = str(prox)
                        return numero
            if isinstance(cell, (int, float)) and idx > 0:
                ant = row[idx - 1]
                if isinstance(ant, str) and "matr" in ant.lower():
                    numero = str(cell)
                    return numero
    print("Número de matrícula não encontrado")
    return ""

def extrair_valor_total_excel(caminho_arquivo):
    wb = openpyxl.load_workbook(caminho_arquivo, data_only=True)
    aba = "SANEAMENTO"
    if aba not in wb.sheetnames:
        print(f"{aba} não encontrado")
        return ""
    ws = wb[aba]
    for row in ws.iter_rows(values_only=True):
        for idx, cell in enumerate(row):
            if isinstance(cell, str) and cell.strip().lower() == "valor total":
                for prox in row[idx+1:]:
                    if prox not in (None, "", "-"):
                        try:
                            valor = float(prox)
                            valor_formatado = f"R$ {valor:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
                        except Exception:
                            valor_formatado = str(prox)
                        print(f"Valor Total encontrado: {valor_formatado}")
                        return valor_formatado
    print("Valor Total não encontrado")
    return ""
    
def extrair_valor_liq_excel(caminho_arquivo):
    wb = openpyxl.load_workbook(caminho_arquivo, data_only=True)
    aba = "LIQUIDAÇÃO"
    ws = wb[aba]

    for row in ws.iter_rows(values_only=True):
        for idx, cell in enumerate(row):
            if isinstance(cell, str) and cell.strip().lower() == "valor de liquidação forçada":
                for prox in row[idx+1:]:
                    if prox not in (None, "", "-"):
                        try:
                            valor = float(prox)
                            valor_formatado = f"R$ {valor:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
                        except Exception:
                            valor_formatado = str(prox)
                        print(f"Valor de Liquidação Forçada encontrado: {valor_formatado}")
                        return valor_formatado

    print("Valor de Liquidação Forçada não encontrado")
    return ""

def extrair_area_total_excel(caminho_arquivo):
    aba = "AREA UTIL"
    try:
        wb = openpyxl.load_workbook(caminho_arquivo, data_only=True)
    except Exception as e:
        print(f"Erro ao abrir a planilha: {e}")
        return ""

    if aba not in wb.sheetnames:
        print(f"Aba '{aba}' não encontrada. Abas disponíveis: {wb.sheetnames}")
        return ""

    ws = wb[aba]

    for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
        for col_idx, cell in enumerate(row):
            if isinstance(cell, str):
                texto = cell.strip().lower()
                if texto == "area total":
                    try:
                        valor = ws.cell(row=row_idx + 2, column=col_idx + 1).value
                        if valor not in (None, "", "-"):
                            try:
                                valor = float(valor)
                                valor_formatado = f"{valor:.4f}".replace(".", ",")
                            except Exception:
                                valor_formatado = str(valor)
                            print(f"✅ Área TOTAL encontrada: {valor_formatado}")
                            return valor_formatado
                        else:
                            print(f"⚠️ Célula abaixo de 'AREA TOTAL' está vazia.")
                    except Exception as e:
                        print(f"❌ Erro ao acessar célula abaixo: {e}")
                        return ""

    print("❌ 'AREA TOTAL' não encontrada na aba.")
    return ""
    
def extrair_area_const_excel(caminho_arquivo):
    wb = openpyxl.load_workbook(caminho_arquivo, data_only=True)
    aba = "AREA UTIL"

    if aba not in wb.sheetnames:
        print(f"Aba '{aba}' não encontrada. Abas disponíveis: {wb.sheetnames}")
        return ""

    ws = wb[aba]

    for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
        for col_idx, cell in enumerate(row):
            if isinstance(cell, str):
                texto = unicodedata.normalize("NFKD", cell).encode("ASCII", "ignore").decode().strip().lower()
                if texto == "area consolidada":
                    try:
                        valor = ws.cell(row=row_idx + 2, column=col_idx + 1).value
                        if valor not in (None, "", "-"):
                            try:
                                valor = float(valor)
                                valor_formatado = f"{valor:.4f}".replace(".", ",")
                            except Exception:
                                valor_formatado = str(valor)
                            print(f"✅ ÁREA CONSOLIDADA encontrada: {valor_formatado}")
                            return valor_formatado
                        else:
                            print(f"⚠️ Célula abaixo de 'ÁREA CONSOLIDADA' está vazia.")
                    except Exception as e:
                        print(f"❌ Erro ao acessar célula abaixo: {e}")
                        return ""

    print("❌ 'ÁREA CONSOLIDADA' não encontrada na aba.")
    return ""

def extrair_porcentagem_reserva_excel(caminho_arquivo):
    wb = openpyxl.load_workbook(caminho_arquivo, data_only=True)
    aba = "AREA UTIL"
    if aba not in wb.sheetnames:
        print(f"{aba} não encontrado")
        return ""
    
    ws = wb[aba]
    
    for row in ws.iter_rows(values_only=True):
        for idx, cell in enumerate(row):
            if isinstance(cell, str):
                texto = unicodedata.normalize("NFKD", cell).encode("ASCII", "ignore").decode().strip().lower()
                if texto == "area de reserva legal":
                    valores_validos = [v for v in row[idx+2:] if v not in (None, "", "-")]
                    if len(valores_validos) >= 3:
                        try:
                            valor = float(valores_validos[2])  # terceiro valor válido
                            valor_formatado = f"{valor:.2f}%"
                        except Exception:
                            valor_formatado = str(valores_validos[2])
                        print(f"✅ Área de Reserva Legal encontrada: {valor_formatado}")
                        return valor_formatado
                    else:
                        print("⚠️ Menos de 3 valores após 'Área de Reserva Legal'")
                        return valor_formatado
    print("⚠️ Procentagem de Reserva Legal não encontrada")
    return ""

def extrair_area_reserva_excel(caminho_arquivo):
    wb = openpyxl.load_workbook(caminho_arquivo, data_only=True)
    aba = "AREA UTIL"
    if aba not in wb.sheetnames:
        print(f"{aba} não encontrado")
        return ""
    ws = wb[aba]
    for row in ws.iter_rows(values_only=True):
        for idx, cell in enumerate(row):
            if isinstance(cell, str) and cell.strip().lower() == "Área De Reserva Legal":
                for prox in row[idx+1:]:
                    if prox not in (None, "", "-"):
                        try:
                            valor = float(prox)
                            valor_formatado = f"{valor:.4f}"
                        except Exception:
                            valor_formatado = str(valor)
                        print(f"Área De Reserva Legal encontrado: {valor_formatado}")
                        return valor_formatado
    print("Área De Reserva Legal não encontrado")
    return ""

def extrair_porcentagem_app_excel(caminho_arquivo):
    wb = openpyxl.load_workbook(caminho_arquivo, data_only=True)
    aba = "AREA UTIL"
    if aba not in wb.sheetnames:
        print(f"{aba} não encontrado")
        return ""
    ws = wb[aba]
    for row in ws.iter_rows(values_only=True):
        for idx, cell in enumerate(row):
            if isinstance(cell, str) and cell.strip().lower() == "Área de Presevarção Permanente":
                for prox in row[idx+2:]:
                    if prox not in (None, "", "-"):
                        try:
                            valor = float(prox)
                            valor_formatado = f"{valor:.2f}%"
                        except Exception:
                            valor_formatado = str(valor)
                        print(f"Área de Presevarção Permanente encontrado: {valor_formatado}")
                        return valor_formatado
    print("Área de Presevarção Permanente não encontrado")
    return ""

def extrair_area_app_excel(caminho_arquivo):
    wb = openpyxl.load_workbook(caminho_arquivo, data_only=True)
    aba = "AREA UTIL"
    if aba not in wb.sheetnames:
        print(f"{aba} não encontrado")
        return ""
    ws = wb[aba]
    for row in ws.iter_rows(values_only=True):
        for idx, cell in enumerate(row):
            if isinstance(cell, str) and cell.strip().lower() == "Área de Presevarção Permanente":
                for prox in row[idx+1:]:
                    if prox not in (None, "", "-"):
                        try:
                            valor = float(prox)
                            valor_formatado = f"{valor:.4f}"
                        except Exception:
                            valor_formatado = str(valor)
                        print(f"Área de Presevarção Permanente encontrado: {valor_formatado}")
                        return valor_formatado
    print("Área de Presevarção Permanente não encontrado")
    return ""