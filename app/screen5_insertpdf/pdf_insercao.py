from kivy.metrics import dp
from kivymd.uix.snackbar import MDSnackbar, MDSnackbarText
from docx.shared import Cm
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION
import os
import time
import re
from modules.tabela_excel_para_word import *
from app.screen5_insertpdf.pdf_extracao import extrair_paginas_como_imagens
from assets.document_style import *
from assets.document_page1 import *
from assets.document_page2 import *
from assets.document_page3 import *
from assets.document_page4 import *
from assets.document_page5 import *
from assets.document_page6 import *
from assets.document_page7 import *
from assets.document_page8 import *

def inserir_pdf_no_word(self, caminho_pdf, placeholder):
    if not hasattr(self, "doc"):
        print("Documento Word não inicializado!")
        return
    else:
        print("Documento iniciado")

    def substituir_em_paragrafos(paragrafos):
        for par in paragrafos:
            texto_completo = ''.join(run.text for run in par.runs)
            if placeholder in texto_completo:
                for run in par.runs:
                    run.text = ""
                for imagem_path in imagens:
                    novo_run = par.add_run()
                    novo_run.add_picture(imagem_path, width=Cm(14))
                return True
        return False

    try:
        imagens = extrair_paginas_como_imagens(caminho_pdf)
        if not imagens:
            MDSnackbar(
                MDSnackbarText(text="Nenhuma imagem encontrada no PDF."),
                y=dp(24)
            ).open()
            return

        encontrado = substituir_em_paragrafos(self.doc.paragraphs)

        if not encontrado:
            for tabela in self.doc.tables:
                for linha in tabela.rows:
                    for celula in linha.cells:
                        if substituir_em_paragrafos(celula.paragraphs):
                            encontrado = True
                            break
                    if encontrado:
                        break
                if encontrado:
                    break

        if encontrado:
            if not hasattr(self, "pdfs_inseridos"):
                self.pdfs_inseridos = set()
            self.pdfs_inseridos.add(placeholder)

    except Exception as e:
        MDSnackbar(
            MDSnackbarText(text=f"Erro ao inserir PDF: {str(e)}"),
            y=dp(24)
        ).open()

def inserir_imagem_no_placeholder(self, placeholder, caminho_imagem):
    for par in self.doc.paragraphs:
        texto_completo = ''.join(run.text for run in par.runs)
        if placeholder in texto_completo:
            for run in par.runs:
                run.text = ""
            if os.path.exists(caminho_imagem):
                novo_run = par.add_run()
                novo_run.add_picture(caminho_imagem, width=Cm(14))
                print(f"✅ Imagem '{caminho_imagem}' inserida no placeholder '{placeholder}'")
            else:
                print(f"❌ Caminho inválido: {caminho_imagem}")
            return
    print(f"❌ Placeholder '{placeholder}' não encontrado no documento.")


def montar_documento(self, doc):
    """
    Monta o documento Word, criando seções por imóvel utilizando
    self.lista_dados_matriculas já normalizada (nenhuma duplicação).
    """
    print(f"📋 Quantidade de matrículas em montar_documento: {len(self.lista_dados_matriculas)}")
    doc = configurar_documento()
    doc.add_page_break()
    print(f"📋 Quantidade de matrículas em montar_documento depois de configurar: {len(self.lista_dados_matriculas)}")

    for i, dados in enumerate(self.lista_dados_matriculas):
        doc = criar_titulo(doc, dados)
        doc = adicionar_linha_fina(doc)
        doc = criar_secao_valor(doc, dados)
        doc = adicionar_linha_fina(doc)
        doc = criar_secao_identificacao(doc)
        doc = adicionar_linha_fina(doc)
        doc = criar_secao_croqui(doc, imagem_path=dados.get("imagem", ""))
        doc = adicionar_linha_fina(doc)
        doc = geometria_terreno(doc, dados)
        doc = adicionar_linha_fina(doc)
        doc = criar_secao_caracteristicas(doc, dados)
        if i < self.qtd_imoveis - 1:
            doc.add_page_break()

    for i, dados in enumerate(self.lista_dados_matriculas):
        doc = adicionar_linha_fina(doc)
        doc = titulo(doc, dados)
        doc = table_geo(doc, dados)
        doc = adicionar_linha_fina(doc)
        doc = tabela_bioma(doc, dados)
        doc = adicionar_linha_fina(doc)
        doc = area_APA(doc, dados)
        doc = adicionar_linha_fina(doc)
        doc = table_passivo_ambiental(doc, dados)
        doc.add_page_break()

    doc = inserir_sumario(doc)
    doc.add_page_break()
    doc = adicionar_espaco(doc)
    doc = texto_solicitante(doc, self.solicitante, self.lista_dados_matriculas)
    doc = adicionar_espaco(doc)
    doc = texto_objetivo(doc, self.lista_dados_matriculas)
    doc = adicionar_espaco(doc)
    doc = texto_finalidade(doc)
    doc = adicionar_espaco(doc)
    doc = texto_proprietario(doc, self.lista_dados_matriculas, self.lista_proponentes)
    doc = adicionar_espaco(doc)
    doc = texto_ressalvas(doc)
    doc = title_imovel(doc)
    doc = localizacao(doc, self.lista_dados_matriculas)
    doc = acesso(doc)
    doc = carac_reg(doc)
    doc.add_page_break()
    doc = desc_imovel(doc, self.lista_dados_matriculas)
    doc.add_page_break()
    doc = declividade(doc, imagem_path=getattr(self, "caminho_declividade", None))
    doc.add_page_break()
    doc = hidrografia(doc, imagem_path=getattr(self, "caminho_hidrografia", None))
    doc.add_page_break()
    doc = pedologia(doc, imagem_path=getattr(self, "caminho_solos", None))
    doc.add_page_break()
    doc = uso_imovel(doc)
    doc = adicionar_espaco(doc)
    doc = benfeitoria(doc)
    doc = diag_mercado(doc)
    doc = metodologia(doc)
    doc.add_page_break()
    doc = metodo_comparativo(doc)
    doc = adicionar_espaco(doc)
    doc = aproveitamento(doc)
    doc = adicionar_espaco(doc)
    doc = especificacao(doc)
    doc = grau_especificacao(doc)
    doc.add_page_break()
    doc = grau_precisao(doc)
    doc = adicionar_espaco(doc)
    doc = grau_precisao2(doc)
    doc.add_page_break()
    doc = resultado(doc, self.lista_dados_matriculas)
    doc.add_page_break()
    doc = encerramento(doc, self.lista_dados_matriculas[0] if self.lista_dados_matriculas else {})
    doc.add_page_break()
    doc = inserir_caixa_texto(doc)
    doc.add_page_break()
    doc = anexos_fotos(doc)
    doc.add_page_break()
    doc = anexo_doc(doc)
    doc = adicionar_espaco(doc)
    doc = anexo_parametros(doc, self.lista_dados_matriculas)

    return doc

def gerar_documento(self):
    #try: 
        self.imagem_marca_dagua = "models/RODAPE.png"
        self.imagem_final = "models/final.png"
        self.img_capa = "models/capa_do_laudo.png"
        self.img_anexo = "models/anexos.png"
        processo = ""

        if hasattr(self, "caminho_car") and self.caminho_car:
            caminho_processo = os.path.dirname(self.caminho_car)
        elif hasattr(self, "caminho_cit") and self.caminho_cit:
            caminho_processo = os.path.dirname(self.caminho_cit)
        else:
            caminho_processo = ""

        if caminho_processo:
            match = re.search(r"processo\s*n[\u00b0\u00ba]?\s*(\d+)", caminho_processo, re.IGNORECASE)
            if match:
                processo = match.group(1)

        nome_sanitizado = re.sub(r'[\\/*?:"<>|]', "_", self.solicitante)
        nome_arquivo = f"LAUDO DE AVALIACAO Nº {processo} {nome_sanitizado}.docx"

        if self.current_path and os.path.exists(self.current_path):
            enviados_dir = os.path.join(self.current_path, "ENVIADOS")
            os.makedirs(enviados_dir, exist_ok=True)
            output_path = os.path.join(enviados_dir, nome_arquivo)
        else:
            output_path = os.path.join(os.getcwd(), "output", nome_arquivo)
            os.makedirs(os.path.dirname(output_path), exist_ok=True)

        output_path = os.path.normpath(output_path)

        doc = configurar_documento()
        doc.save(output_path)
        time.sleep(1)

        texto_capa = "LAUDO DE AVALIAÇÃO Nº #NPROCESSO,\n #DATA_ATUAL, Palmas TO"
        substituicoes = {
            "#NPROCESSO": processo,
            "#DATA_ATUAL": self.data_atual
        }

        doc = Document(output_path)
        self.doc = montar_documento(self, doc)

        if hasattr(self, "caminho_declividade"):
            inserir_imagem_no_placeholder(self, "#IMAGEM_DECLIVIDADE", self.caminho_declividade)
        if hasattr(self, "caminho_hidrografia"):
            inserir_imagem_no_placeholder(self, "#IMAGEM_HIDROGRAFIA", self.caminho_hidrografia)
        if hasattr(self, "caminho_rotas"):
            inserir_imagem_no_placeholder(self, "#IMAGEM_ACESSO", self.caminho_rotas)
        if hasattr(self, "caminho_solos"):
            inserir_imagem_no_placeholder(self, "#IMAGEM_SOLOS", self.caminho_solos)

        substituicoes_base = {
            "#TRATAMENTO": self.tratamento,
            "#PROPONENTE": self.nome,
            "#CPF_PROPONENTE": self.cpf,
            "#DATA_ATUAL": self.data_atual,
            "#CIVIL": self.civil,
            "#CIDADE_I": self.municipio,
            "#ESTADO_I": self.estado,
            "#SOLICITANTE": self.solicitante,
            "#DESCRICAO_IMOVEL": self.descricao_imovel,
            "#REGIAO_CIDADE": self.descricao_cidade,
            "#ATIVIDADE_IMOVEL": self.atividade_imovel,
            "#REGIAO_IMOVEL": self.regiao_imovel,
            "#DECLIVIDADE_I": self.declividade,
            "#HIDROGRAFIA_I": self.hidrografia,
            "#TIPO_SOLO": self.resumo_solo,
            "#DESCRICAO_SOLO": self.texto_solos,
            "#ROTA_ACESSO": self.rotas,
            "#NPROCESSO": processo,
        }

        def substituir_texto(substituicoes):
            def substituir_em_runs(par):
                for run in par.runs:
                    for chave, valor in substituicoes.items():
                        if chave in run.text:
                            run.text = run.text.replace(chave, valor)

            def substituir_em_paragrafos(paragrafos):
                for par in paragrafos:
                    substituir_em_runs(par)

            def substituir_em_tabela(tabela):
                for linha in tabela.rows:
                    for celula in linha.cells:
                        substituir_em_paragrafos(celula.paragraphs)
                        for tabela_interna in celula.tables:
                            substituir_em_tabela(tabela_interna)

            substituir_em_paragrafos(self.doc.paragraphs)
            for tabela in self.doc.tables:
                substituir_em_tabela(tabela)
            for section in self.doc.sections:
                substituir_em_paragrafos(section.header.paragraphs)
            for shape in self.doc.inline_shapes:
                if shape._inline.graphic.graphicData.uri.endswith("/wordprocessingShape"):
                    for box in shape._inline.graphic.graphicData.xpath(".//w:txbxContent"):
                        for par_el in box.iter(qn('w:p')):
                            for r in par_el.iter(qn('w:t')):
                                if r.text:
                                    for chave, valor in substituicoes.items():
                                        if chave in r.text:
                                            r.text = r.text.replace(chave, valor)

        for i, dados in enumerate(self.lista_dados_matriculas):
            substituicoes = substituicoes_base.copy()
            substituicoes.update({
                "#NOME_IMOVEL": dados.get("nome_imovel", ""),
                "#LATITUDE": dados.get("latitude", ""),
                "#LONGITUDE": dados.get("longitude", ""),
                "#NMATRICULA": dados.get("matricula", ""),
                "#VALOR_TOTAL": dados.get("valor_total", ""),
                "#VALOR_LIQUIDO": dados.get("valor_liq", ""),
            })
            substituir_texto({k: str(v) for k, v in substituicoes.items()})

        if self.caminho_car:
            inserir_pdf_no_word(self, self.caminho_car, "#SUBSTITUIR_CAR")
        if self.caminho_cit:
            inserir_pdf_no_word(self, self.caminho_cit, "#SUBSTITUIR_CIT")

        self.doc.save(output_path)

        unique_matriculas = set(d.get("matricula") for d in self.lista_dados_matriculas)
        print("🧾 Total de matrículas carregadas:", len(unique_matriculas))

        matriculas_processadas = set()
        for i, dados in enumerate(self.lista_dados_matriculas):
            matricula = dados.get("matricula")
            if matricula in matriculas_processadas:
                continue
            matriculas_processadas.add(matricula)

            caminho_excel = dados.get("planilha")
            if caminho_excel and os.path.exists(caminho_excel):
                marcador = f"[INSERIR_TABELA_{i}_AQUI]"
                marcador_homog = f"[INSERIR_HOMOG_{i}_AQUI]"
                marcador_saneamento = f"[INSERIR_SANEAMENTO_{i}_AQUI]"
                marcador_quadro = f"[INSERIR_QUADRO_{i+1:02d}_AQUI]"
                marcador_liq = f"[INSERIR_LIQUIDACAO_{i}_AQUI]"
                marcador_valores = f"[INSERIR_VALORES_{i}_AQUI]"
                print(f"📌 Inserindo tabelas para matrícula índice {i}")
                inserir_tabela_excel_no_word(output_path, caminho_excel, marcador_personalizado=marcador)
                inserir_tabela_benfeitoria_no_word(output_path, caminho_excel)
                inserir_tabela_depreciacao_no_word(output_path, caminho_excel)
                inserir_tabela_classe_no_word(output_path, caminho_excel)
                inserir_tabelas_amostras_auto(output_path, caminho_excel, indice_matricula=i)
                inserir_tabela_situacao_no_word(output_path, caminho_excel)
                inserir_tabela_quadro_no_word(output_path, caminho_excel, marcador_quadro)
                inserir_tabela_homog_no_word(output_path, caminho_excel, marcador_homog)
                inserir_tabela_saneamento_no_word(output_path, caminho_excel, marcador_saneamento)
                inserir_tabela_liquidacao_no_word(output_path, caminho_excel, marcador_liq)
                inserir_tabela_valores_no_word(output_path, caminho_excel, marcador_valores)

        inserir_e_atualizar_sumario_no_bookmark(output_path, bookmark_name="SUMARIO")

        if hasattr(self, "imagem_marca_dagua"):  
            imagens_fundo(output_path, self.imagem_marca_dagua)
        if hasattr(self, "imagem_final"):
            inserir_imagem_ultima_pagina(output_path, self.imagem_final)
        if hasattr(self, "img_anexo"):
            inserir_marcadagua_so_na_secao(output_path, self.img_anexo, marcador='#CAIXATEXTO#')
        if hasattr(self, "img_capa"):
            inserir_imagem_capa_atras_texto(output_path, self.img_capa)

        inserir_caixa_texto_primeira_pagina(output_path, texto_capa, substituicoes=substituicoes)

        self.word_app = win32com.client.Dispatch("Word.Application")
        MDSnackbar(
            MDSnackbarText(text="\u2705Documento gerado com sucesso!"),
            y=dp(24)
        ).open()

"""    except Exception as e:
        print(f"❌ Erro ao gerar documento: {e}")
        try:
            word = win32com.client.GetActiveObject("Word.Application")
            word.Quit()
            if hasattr(self, "word_app"):
                try:
                    self.word_app.Quit()
                    print("✅ Word encerrado com sucesso")
                except Exception as quit_err:
                    print(f"❌ Erro ao encerrar o Word: {quit_err}")
        except Exception as close_err:
            print(f"⚠️ Erro ao tentar fechar o Word: {close_err}")
        MDSnackbar(
            MDSnackbarText(text=f"Erro: {str(e)}"),
            y=dp(24)
        ).open()"""
