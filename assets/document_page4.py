from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import win32com.client

def inserir_sumario(doc):
    paragraph = doc.add_paragraph()
    run = paragraph.add_run("SUMÁRIO")
    tag_start = OxmlElement('w:bookmarkStart')
    tag_start.set(qn('w:id'), '1')
    tag_start.set(qn('w:name'), 'SUMARIO')
    paragraph._p.append(tag_start)
    paragraph.add_run()
    tag_end = OxmlElement('w:bookmarkEnd')
    tag_end.set(qn('w:id'), '1')
    paragraph._p.append(tag_end)

    paragraph = doc.add_paragraph()
    runpar = paragraph.add_run("SUMÁRIO")
    paragraph.runs[0].bold = True
    return doc

def inserir_e_atualizar_sumario_no_bookmark(docx_path, bookmark_name="SUMARIO"):
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    doc = word.Documents.Open(docx_path)

    if hasattr(doc, "Bookmarks") and doc.Bookmarks.Exists(bookmark_name):
        rng = doc.Bookmarks(bookmark_name).Range
        doc.TablesOfContents.Add(
            Range=rng,
            RightAlignPageNumbers=True,
            UseHeadingStyles=True,
            UpperHeadingLevel=1,
            LowerHeadingLevel=3,
            IncludePageNumbers=True,
            AddedStyles="",
            UseHyperlinks=True,
            HidePageNumbersInWeb=True,
            UseOutlineLevels=True
        )
        doc.TablesOfContents(1).Update()
        print("✅ Sumário inserido e atualizado!")
    else:
        print(f"❌ Bookmark '{bookmark_name}' não encontrado.")

    doc.Save()
    doc.Close(False)
    #word.Quit()