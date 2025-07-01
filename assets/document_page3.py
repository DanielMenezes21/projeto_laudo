from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import parse_xml
from docx.table import Table
from docx.oxml.ns import nsdecls
from docx.shared import Twips
from assets.document_page2 import LARGURA, configurar_documento, adicionar_linha_fina

def titulo(doc, dados):
    table = doc.add_table(rows=1, cols=1)
    table.allow_autofit = False
    usable_width = LARGURA
    table.width = usable_width
    table.style = 'Table Grid'

    cell = table.cell(0, 0)

    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="4EA65D"/>')
    cell._tc.get_or_add_tcPr().append(shading)

    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    cell._tc.get_or_add_tcPr().append(borders)

    matricula = dados.get("matricula", "")
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"Parecer Jurídico-ambiental da matrícula {matricula}")
    run.bold = True
    run.font.size = Pt(12)

    return doc

def table_geo(doc):
    table = doc.add_table(rows=12, cols=10)
    table.allow_autofit = False
    usable_width = LARGURA
    table.width = usable_width
    table.style = 'Table Grid'

    table.rows[0].height = Cm(0.5)
    table.rows[1].height = Cm(0.5)  
    table.rows[2].height = Cm(0.1)
    table.rows[3].height = Cm(0.5)
    table.rows[4].height = Cm(0.1)
    table.rows[5].height = Cm(0.5)
    table.rows[6].height = Cm(0.9)
    table.rows[7].height = Cm(0.5)
    table.rows[8].height = Cm(0.5)
    table.rows[9].height = Cm(0.5)
    table.rows[10].height = Cm(1.2)
    table.rows[11].height = Cm(0.1)

    linha1 = table.rows[0].cells[0]
    linha1.merge(table.rows[0].cells[9])
    par = linha1.paragraphs[0]
    par.clear()
    run = par.add_run("Possui Georreferenciamento?")
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run.font.size = Pt(12)
    run.bold = True
    tcPr = linha1._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell2_0 = table.rows[1].cells[0]
    par = cell2_0.paragraphs[0]
    par.clear()
    run = par.add_run("Sim")
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cell2_0.width = Cm(4)
    run.font.size = Pt(12)
    tcPr = cell2_0._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell2_1 = table.rows[1].cells[1]
    par = cell2_1.paragraphs[0]
    par.clear()
    run = par.add_run(" ")
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cell2_1.width = Cm(0.4)
    run.font.size = Pt(12)
    tcPr = cell2_1._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell2_2 = table.rows[1].cells[2]
    par = cell2_2.paragraphs[0]
    par.clear()
    run = par.add_run("Não")
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run.font.size = Pt(12)
    tcPr = cell2_2._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell2_3 = table.rows[1].cells[3]
    cell2_3.text = " "
    par = cell2_3.paragraphs[0]
    par.clear()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cell2_3.width = Cm(0.4)
    tcPr = cell2_3._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell2_4 = table.rows[1].cells[4]
    cell2_5 = table.rows[1].cells[5]
    cell2_4.merge(cell2_5)
    par = cell2_4.paragraphs[0]
    par.clear()
    run = par.add_run("Código Geo")
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cell2_4.width = Cm(3.5)
    tcPr = cell2_4._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell2_6 = table.rows[1].cells[6]
    cell2_6.merge(table.rows[1].cells[8])
    par = cell2_6.paragraphs[0]
    par.clear()
    run = par.add_run(" ")
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cell2_6.width = Cm(3.5)
    run.font.size = Pt(12)
    tcPr = cell2_6._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)
    cell2_7 = table.rows[1].cells[9]
    cell2_7.text = " "
    cell2_7.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cell2_7.paragraphs[0].runs[0].font.size = Pt(12)
    tcPr = cell2_7._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    linha3 = table.rows[2].cells[0]
    linha3.merge(table.rows[2].cells[9])
    paragraph = linha3.paragraphs[0]
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT 
    run = paragraph.add_run(" ")
    run.font.size = Pt(1)
    tcPr = linha3._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    linha4 = table.rows[3].cells[0]
    par = linha4.paragraphs[0]
    par.clear()
    run = par.add_run("CAR")
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run.font.size = Pt(12)
    tcPr = linha4._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell4_1 = table.rows[3].cells[1]
    cell4_1.merge(table.rows[3].cells[7])
    par = cell4_1.paragraphs[0]
    par.clear()
    run = par.add_run(" ")
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cell4_1.width = Cm(5.5)
    tcPr = cell4_1._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell4_2 = table.rows[3].cells[8]
    cell4_2.merge(table.rows[3].cells[9])
    par = cell4_2.paragraphs[0]
    par.clear()
    run = par.add_run(" ")
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cell4_2.width = Cm(3.5)
    tcPr = cell4_2._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    linha5 = table.rows[4].cells[0]
    linha5.merge(table.rows[4].cells[9])
    paragraph = linha5.paragraphs[0]
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT 
    run = paragraph.add_run("  ")
    run.font.size = Pt(3)
    tcPr = linha5._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    linha6 = table.rows[5].cells[0]
    linha6.merge(table.rows[5].cells[1])
    par = cell2_2.paragraphs[0]
    par.clear()
    run = par.add_run("Não")
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run.font.size = Pt(12)
    tcPr = linha6._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)
    cell6_2 = table.rows[5].cells[2]
    cell6_2.merge(table.rows[5].cells[4])
    par = cell6_2.paragraphs[0]
    par.clear()
    run = par.add_run(" ")
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run.font.size = Pt(12)
    tcPr = cell6_2._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell6_5 = table.rows[5].cells[5]
    par = cell6_5.paragraphs[0]
    par.clear()
    run = par.add_run("CPF")
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run.font.size = Pt(12)
    run.bold = True
    tcPr = cell6_5._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell6_6 = table.rows[5].cells[6]
    cell6_6.merge(table.rows[5].cells[8])
    par = cell6_6.paragraphs[0]
    par.clear()
    run = par.add_run(" ")
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run.font.size = Pt(12)
    tcPr = cell6_6._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell6_9 = table.rows[5].cells[9]
    par = cell6_9.paragraphs[0]
    par.clear()
    run = par.add_run(" ")
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run.font.size = Pt(12)
    tcPr = cell6_9._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    linha7 = table.rows[6].cells[0]
    linha7.merge(table.rows[6].cells[9])
    paragraph = linha7.paragraphs[0]
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run("Possui Alienação Fiduciária?")
    run.font.size = Pt(12)
    run.bold = True
    tcPr = linha7._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell7_1 = table.rows[7].cells[0]
    par = cell7_1.paragraphs[0]
    par.clear()
    run = par.add_run("Sim")
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cell7_1.width = Cm(4)
    run.font.size = Pt(12)
    tcPr = cell7_1._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell7_2 = table.rows[7].cells[1]
    par = cell7_2.paragraphs[0]
    par.clear()
    run = par.add_run(" ")
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cell7_2.width = Cm(0.4)
    run.font.size = Pt(12)
    tcPr = cell7_2._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)
    cell7_3 = table.rows[7].cells[2]
    par = cell7_3.paragraphs[0]
    par.clear()
    run = par.add_run("Não")
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run.font.size = Pt(12)
    tcPr = cell7_3._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell7_4 = table.rows[7].cells[3]
    par = cell7_4.paragraphs[0]
    par.clear()
    run = par.add_run(" ")
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run.font.size = Pt(12)
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell7_5 = table.rows[7].cells[4]
    cell7_5.merge(table.rows[7].cells[9])
    par = cell7_5.paragraphs[0]
    par.clear()
    run = par.add_run(" ")
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run.font.size = Pt(12)
    tcPr = cell7_5._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    linha8 = table.rows[8].cells[0]
    linha8.merge(table.rows[8].cells[4])
    par = linha8.paragraphs[0]
    par.clear()
    run = par.add_run(" ")
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run.font.size = Pt(12)
    tcPr = linha8._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="nil"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell8_5 = table.rows[8].cells[5]
    par = cell8_5.paragraphs[0]
    par.clear()
    run = par.add_run("Averbação")
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run.font.size = Pt(12)
    run.bold = True
    tcPr = cell8_5._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="nil"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell8_6 = table.rows[8].cells[6]
    cell8_6.merge(table.rows[8].cells[8])
    par = cell8_6.paragraphs[0]
    par.clear()
    run = par.add_run(" ")
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run.font.size = Pt(12)
    tcPr = cell8_6._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell8_9 = table.rows[8].cells[9]
    par = cell8_9.paragraphs[0]
    par.clear()
    run = par.add_run(" ")
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run.font.size = Pt(12)
    tcPr = cell8_9._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    linha9 = table.rows[9].cells[0]
    linha9.merge(table.rows[9].cells[9])
    paragraph = linha9.paragraphs[0]
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run("Outras observações documentadas")
    run.font.size = Pt(12)
    run.bold = True
    tcPr = linha9._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    linha10 = table.rows[10].cells[0]
    linha10.merge(table.rows[10].cells[5])
    paragraph = linha10.paragraphs[0]
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(" ")
    run.font.size = Pt(12)
    run.bold = True
    run.space_after = Twips(0)
    run.space_before = Twips(0)
    tcPr = linha10._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell10_6 = table.rows[10].cells[6]
    cell10_6.merge(table.rows[10].cells[8])
    par = cell10_6.paragraphs[0]
    par.clear()
    run = par.add_run(" ")
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run.font.size = Pt(12)
    tcPr = cell10_6._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell10_9 = table.rows[10].cells[9]
    par = cell10_9.paragraphs[0]
    par.clear()
    run = par.add_run(" ")
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run.font.size = Pt(12)
    tcPr = cell10_9._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    linha11 = table.rows[11].cells[0]
    linha11.merge(table.rows[11].cells[9])
    paragraph = linha11.paragraphs[0]
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(" ")
    run.font.size = Pt(1)
    run.space_after = Twips(0)
    run.space_before = Twips(0)
    tcPr = linha11._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    return doc

def tabela_bioma(doc):
    """
    Cria uma tabela no documento Word com informações sobre biomas.
    """
    table = doc.add_table(rows=7, cols=10)
    table.allow_autofit = False
    usable_width = LARGURA
    table.width = usable_width
    table.style = 'Table Grid'

    table.rows[0].height = Cm(0.5)
    table.rows[1].height = Cm(0.5)  
    table.rows[2].height = Cm(0.1)
    table.rows[3].height = Cm(0.5)
    table.rows[4].height = Cm(0.1)
    table.rows[5].height = Cm(0.5)
    table.rows[6].height = Cm(0.1)

    title_row = table.rows[0].cells[0]
    title_row.merge(table.rows[0].cells[9])
    title_row.text = "Bioma no qual a propriedade está inserida"
    title_row.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_row.paragraphs[0].runs[0].font.size = Pt(12)
    title_row.paragraphs[0].runs[0].bold = True
    tcPr = title_row._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell1_0 = table.rows[1].cells[0]
    par = cell1_0.paragraphs[0]
    par.clear()
    run = par.add_run("Amazônia")
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run.font.size = Pt(12)
    cell1_0.width = Cm(4)
    tcPr = cell1_0._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell1_1 = table.rows[1].cells[1]
    par = cell1_1.paragraphs[0]
    par.clear()
    run = par.add_run(" ")
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run.font.size = Pt(12)
    cell1_1.width = Cm(0.4)
    tcPr = cell1_1._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>' 
        '<w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell1_2 = table.rows[1].cells[2]
    cell1_2.merge(table.rows[1].cells[5])
    par = cell1_2.paragraphs[0]
    par.clear()
    run = par.add_run("Pampa")
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run.font.size = Pt(12)
    tcPr = cell1_2._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell1_3 = table.rows[1].cells[6]
    par = cell1_3.paragraphs[0]
    par.clear()
    run = par.add_run(" ")
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run.font.size = Pt(12)
    cell1_3.width = Cm(0.4)
    tcPr = cell1_3._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell1_4 = table.rows[1].cells[7]
    cell1_4.merge(table.rows[1].cells[9])
    par = cell1_4.paragraphs[0]
    par.clear()
    run = par.add_run(" ")
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run.font.size = Pt(12)
    tcPr = cell1_4._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    linha3 = table.rows[2].cells[0]
    linha3.merge(table.rows[2].cells[9])
    paragraph = linha3.paragraphs[0]
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("|")
    run.font.size = Pt(1)
    tcPr = linha3._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    linha4 = table.rows[3].cells[0]
    par = linha4.paragraphs[0]
    par.clear()
    run = par.add_run("Cerrado")
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run.font.size = Pt(12)
    tcPr = linha4._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell2_1 = table.rows[3].cells[1]
    par = cell2_1.paragraphs[0]
    par.clear()
    run = par.add_run(" ")
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run.font.size = Pt(12)
    tcPr = cell2_1._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell4_2 = table.rows[3].cells[2]
    cell4_2.merge(table.rows[3].cells[5])
    par = cell4_2.paragraphs[0]
    par.clear()
    run = par.add_run("Pantanal")
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run.font.size = Pt(12)
    tcPr = cell4_2._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell4_4 = table.rows[3].cells[6]
    par = cell4_4.paragraphs[0]
    par.clear()
    run = par.add_run(" ")
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run.font.size = Pt(12)
    tcPr = cell4_4._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell4_5 = table.rows[3].cells[7]
    cell4_5.merge(table.rows[3].cells[9])
    par = cell4_5.paragraphs[0]
    par.clear()
    run = par.add_run(" ")
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run.font.size = Pt(12)
    tcPr = cell4_5._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    linha5 = table.rows[4].cells[0]
    linha5.merge(table.rows[4].cells[9])
    paragraph = linha5.paragraphs[0]
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run('|')
    run.font.size = Pt(1)
    tcPr = linha5._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell6_1 = table.rows[5].cells[0]
    paragraph = cell6_1.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Caatinga")
    run.font.size = Pt(12)
    tcPr = cell6_1._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell6_2 = table.rows[5].cells[1]
    paragraph = cell6_2.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(" ")
    run.font.size = Pt(12)
    tcPr = cell6_2._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell6_3 = table.rows[5].cells[2]
    cell6_3.merge(table.rows[5].cells[5])
    par = cell6_3.paragraphs[0]
    par.clear()
    run = par.add_run("Mata Atlântica")
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run.font.size = Pt(12)
    tcPr = cell6_3._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell6_4 = table.rows[5].cells[6]
    paragraph = cell6_4.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(" ")
    run.font.size = Pt(12)
    tcPr = cell6_4._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell6_5 = table.rows[5].cells[7]
    cell6_5.merge(table.rows[5].cells[9])
    par = cell6_5.paragraphs[0]
    par.clear()
    run = par.add_run(" ")
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run.font.size = Pt(12)
    tcPr = cell6_5._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    linha7 = table.rows[6].cells[0]
    linha7.merge(table.rows[6].cells[9])
    paragraph = linha7.paragraphs[0]
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run('|')
    run.font.size = Pt(1)
    tcPr = linha7._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    return doc

def area_APA(doc):
    """pequena tabela sobre a propriedade e a Área de Proteção Ambiental"""
    table = doc.add_table(rows=3, cols=10)
    table.allow_autofit = False
    usable_width = LARGURA
    table.width = usable_width
    table.style = 'Table Grid'

    table.rows[0].height = Cm(0.5)
    table.rows[1].height = Cm(0.5)  
    table.rows[2].height = Cm(0.1)

    title_row = table.rows[0].cells[0]
    title_row.merge(table.rows[0].cells[9])
    par = title_row.paragraphs[0]
    par.clear()
    run = par.add_run("O imóvel está inserido em área de proteçao ambiental - APA?")
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run.font.size = Pt(12)
    run.bold = True
    tcPr = title_row._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell1_1 = table.rows[1].cells[0]
    par = cell1_1.paragraphs[0]
    par.clear()
    run = par.add_run("sim")
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run.font.size = Pt(12)
    tcPr = cell1_1._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell1_2 = table.rows[1].cells[1]
    par = cell1_1.paragraphs[0]
    par.clear()
    run = par.add_run("sim")
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run.font.size = Pt(12)
    run.bold = True
    cell1_2.width = Cm(0.4)
    tcPr = cell1_2._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell1_3 = table.rows[1].cells[2]
    par = cell1_3.paragraphs[0]
    par.clear()
    run = par.add_run("não")
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run.font.size = Pt(12)
    cell1_3.width = Cm(1.4)
    tcPr = cell1_3._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell1_4 = table.rows[1].cells[3]
    par = cell1_1.paragraphs[0]
    par.clear()
    run = par.add_run(" ")
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run.font.size = Pt(12)
    cell1_4.width = Cm(0.5)
    tcPr = cell1_4._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell1_5 = table.rows[1].cells[4]
    par = cell1_5.paragraphs[0]
    par.clear()
    run = par.add_run("Qual?")
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run.font.size = Pt(12)
    tcPr = cell1_5._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell1_6 = table.rows[1].cells[5]
    cell1_6.merge(table.rows[1].cells[7])
    par = cell1_6.paragraphs[0]
    par.clear()
    run = par.add_run(" ")
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run.font.size = Pt(12)
    tcPr = cell1_6._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell1_7 = table.rows[1].cells[8]
    cell1_7.merge(table.rows[1].cells[9])
    par = cell1_7.paragraphs[0]
    par.clear()
    run = par.add_run(" ")
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run.font.size = Pt(12)
    tcPr = cell1_7._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell2_1 = table.rows[2].cells[0]
    cell2_1.merge(table.rows[2].cells[9])
    paragraph = cell2_1.paragraphs[0]
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run('|')
    run.font.size = Pt(1)
    tcPr = cell2_1._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    return doc

def table_passivo_ambiental(doc):
    '''Tabela de passivo ambiental'''
    table = doc.add_table(rows=11, cols=11)
    table.autofit = False
    usable_width = LARGURA
    table.width = usable_width
    table.style = 'Table Grid'

    table.rows[0].height = Cm(0.5)
    table.rows[1].height = Cm(0.5) 
    table.rows[2].height = Cm(0.15)
    table.rows[3].height = Cm(0.5)
    table.rows[4].height = Cm(0.1)
    table.rows[5].height = Cm(0.5)
    table.rows[6].height = Cm(0.1)
    table.rows[7].height = Cm(0.5)
    table.rows[8].height = Cm(0.1)
    table.rows[9].height = Cm(0.5)
    table.rows[10].height = Cm(0.1)

    title_row = table.rows[0].cells[0]
    title_row.merge(table.rows[0].cells[10])
    par = title_row.paragraphs[0]
    par.clear()
    run = par.add_run("Possui Passivo-Ambiental?")
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run.font.size = Pt(12)
    run.bold = True
    tcPr = title_row._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell1_1 = table.rows[1].cells[0]
    par = cell1_1.paragraphs[0]
    par.clear()
    run = par.add_run("Sim")
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run.font.size = Pt(12)
    tcPr = cell1_1._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell1_2 = table.rows[1].cells[1]
    par = cell1_2.paragraphs[0]
    par.clear()
    run = par.add_run(" ")
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run.font.size = Pt(12)
    cell1_2.width = Cm(0.3)
    tcPr = cell1_2._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell1_3 = table.rows[1].cells[2]
    cell1_3.merge(table.rows[1].cells[3])
    par = cell1_3.paragraphs[0]
    par.clear()
    run = par.add_run("não")
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run.font.size = Pt(12)
    tcPr = cell1_3._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell1_4 = table.rows[1].cells[4]
    par = cell1_4.paragraphs[0]
    par.clear()
    run = par.add_run(" ")
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run.font.size = Pt(12)
    cell1_4.width = Cm(0.3)
    tcPr = cell1_4._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell1_5 = table.rows[1].cells[5]
    cell1_5.merge(table.rows[1].cells[10])
    par = cell1_5.paragraphs[0]
    par.clear()
    run = par.add_run(" ")
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run.font.size = Pt(12)
    tcPr = cell1_5._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell2_1 = table.rows[2].cells[0]
    cell2_1.merge(table.rows[2].cells[10])
    paragraph = cell2_1.paragraphs[0]
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run('|')
    run.font.size = Pt(1)
    tcPr = cell2_1._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell3_1 = table.rows[3].cells[0]
    cell3_1.merge(table.rows[3].cells[1])
    par = cell1_1.paragraphs[0]
    par.clear()
    run = par.add_run("Embargo")
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run.font.size = Pt(12)
    tcPr = cell3_1._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell3_2 = table.rows[3].cells[2]
    par = cell1_1.paragraphs[0]
    par.clear()
    run = par.add_run(" ")
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run.font.size = Pt(12)
    tcPr = cell3_2._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell3_3 = table.rows[3].cells[3]
    cell3_3.merge(table.rows[3].cells[10])
    tcPr = cell3_3._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    linha4 = table.rows[4].cells[0]
    linha4.merge(table.rows[4].cells[10])
    paragraph = linha4.paragraphs[0]
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run('|')
    run.font.size = Pt(1)
    tcPr = linha4._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell4_1 = table.rows[5].cells[0]
    cell4_1.merge(table.rows[5].cells[1])
    par = cell4_1.paragraphs[0]
    par.clear()
    run = par.add_run("Deficit de Reserva Legal")
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run.font.size = Pt(12)
    cell4_1.width = Pt(5)
    tcPr = cell4_1._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell4_2 = table.rows[5].cells[2]
    cell4_2.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cell4_2.width = Cm(0.3)
    tcPr = cell4_2._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell4_3 = table.rows[5].cells[3]
    cell4_3.merge(table.rows[5].cells[10])
    tcPr = cell4_3._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    linha5 = table.rows[6].cells[0]
    linha5.merge(table.rows[6].cells[10])
    paragraph = linha5.paragraphs[0]
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run('|')
    run.font.size = Pt(1)
    tcPr = linha5._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell6_1 = table.rows[7].cells[0]
    cell6_1.merge(table.rows[7].cells[1])
    par = cell6_1.paragraphs[0]
    par.clear()
    run = par.add_run("Alerta MapBiomas")
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run.font.size = Pt(12)
    cell6_1.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    tcPr = cell6_1._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell6_2 = table.rows[7].cells[2]
    cell6_2.width = Cm(0.4)
    tcPr = cell6_2._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell6_3 = table.rows[7].cells[3]
    cell6_3.merge(table.rows[7].cells[10])
    tcPr = cell6_3._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    linha7 = table.rows[8].cells[0]
    linha7.merge(table.rows[8].cells[10])
    paragraph = linha7.paragraphs[0]
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run('|')
    run.font.size = Pt(1)
    tcPr = linha7._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell9_1 = table.rows[9].cells[0]
    cell9_1.merge(table.rows[9].cells[1])
    par = cell9_1.paragraphs[0]
    par.clear()
    run = par.add_run("Detalhamento")
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run.font.size = Pt(12)
    tcPr = cell9_1._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell9_2 = table.rows[9].cells[2]
    cell9_2.merge(table.rows[9].cells[9])
    cell9_2.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
    tcPr = cell9_2._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell9_3 = table.rows[9].cells[10]
    tcPr = cell9_3._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    linha10 = table.rows[10].cells[0]
    linha10.merge(table.rows[10].cells[10])
    paragraph = linha10.paragraphs[0]
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run('|')
    run.font.size = Pt(1)
    tcPr = linha10._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    return doc

def campo_assinatura(doc):
    for _ in range(6):
        par = doc.add_paragraph()
        runpar = par.add_run("")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("_________________________________________________________________")
    run.font.size = Pt(12)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = p2.add_run("ASSINATURA")
    run2.font.size = Pt(14)

    return doc
