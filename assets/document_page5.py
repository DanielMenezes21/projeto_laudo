from docx import Document
from docx.shared import Pt, Cm
from docx.shared import RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Twips
from collections import defaultdict

def adicionar_espaco(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)  
    p.paragraph_format.space_after = Pt(0)   
    run = p.add_run(" ")
    run.font.name = 'Cambria'
    run.font.size = Pt(8) 

    return doc

def gerar_texto_proprietarios_completo(lista_matriculas, lista_proponentes):
    grupos = []

    for item in lista_matriculas:
        proprietarios_nomes = [p.strip() for p in item.get("proprietario", "").split(",") if p.strip()]
        matricula = item.get("matricula", "").strip()
        imovel = item.get("nome_imovel", "").strip()

        for nome in proprietarios_nomes:
            prop_info = next((p for p in lista_proponentes if p["nome"].strip().lower() == nome.strip().lower()), None)

            grupos.append({
                "nome": nome,
                "cpf": prop_info.get("cpf", "") if prop_info else "",
                "civil": prop_info.get("civil", "") if prop_info else "",
                "tratamento": prop_info.get("tratamento", "") if prop_info else "",
                "matricula": matricula,
                "imovel": imovel
            })

    matriculas_dict = defaultdict(list)
    for g in grupos:
        chave = (g['matricula'], g['imovel'])
        matriculas_dict[chave].append(g)

    paragrafos = []
    for (matricula, imovel), proprietarios in matriculas_dict.items():
        if len(proprietarios) == 1:
            p = proprietarios[0]
            texto = (
                f"        Em conformidade com o exposto na matrícula de nº {matricula}, "
                f"o {p['tratamento']} {p['nome']}, {p['civil']}, inscrito sob o CPF nº {p['cpf']}, "
                f"é o proprietário do imóvel rural denominado {imovel}."
            )
        else:
            descricoes = [
                f"{p['tratamento']} {p['nome']}, {p['civil']}, inscrito sob o CPF nº {p['cpf']}"
                for p in proprietarios
            ]
            texto = (
                f"        Em conformidade com o exposto na matrícula de nº {matricula}, "
                f"{formatar_lista_com_e(descricoes)}, são os proprietários do imóvel rural denominado {imovel}."
            )
        paragrafos.append(texto)

    return "\n\n".join(paragrafos)

def formatar_lista_com_e(itens):
    if len(itens) == 1:
        return itens[0]
    elif len(itens) == 2:
        return f"{itens[0]} e {itens[1]}"
    else:
        return ", ".join(itens[:-1]) + f" e {itens[-1]}"
    
def texto_solicitante(doc, solicitante, lista_matriculas):
    heading = doc.add_paragraph(style='Heading1')
    run = heading.add_run("1 - SOLICITANTE")
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.font.name = "Cambria"
    par = doc.add_paragraph()
    runpar = par.add_run("")

    nomes_imoveis = sorted(set([m.get("nome_imovel", "").strip() for m in lista_matriculas if m.get("nome_imovel")]))

    if len(nomes_imoveis) == 1:
        texto = f"        Fomos solicitados pelo {solicitante}, para avaliar um imóvel rural, denominado {nomes_imoveis[0]}, localizado em #CIDADE_I - #ESTADO_I."
    else:
        imoveis = ", ".join(nomes_imoveis[:-1]) + f" e {nomes_imoveis[-1]}"
        texto = f"        Fomos solicitados pelo {solicitante}, para avaliar os imóveis rurais, denominados {imoveis}, localizados em #CIDADE_I - #ESTADO_I."

    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    runpar2 = par.add_run(texto)
    return doc

def texto_objetivo(doc, lista_matriculas):
    heading = doc.add_paragraph( style='Heading1')
    run = heading.add_run("2 - OBJETIVO")
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.font.name = "Cambria"
    par = doc.add_paragraph()
    runpar = par.add_run("")

    nomes_imoveis = sorted(set([m.get("nome_imovel", "").strip() for m in lista_matriculas if m.get("nome_imovel")]))

    if len(nomes_imoveis) == 1:
        texto = f"        O objetivo dessa peça técnica é aferir os valores de mercado e de liquidação forçada por meio do método comparativo de dados de mercado, referente ao imóvel {nomes_imoveis[0]}, localizado em #CIDADE_I - #ESTADO_I."
    else:
        imoveis = ", ".join(nomes_imoveis[:-1]) + f" e {nomes_imoveis[-1]}"
        texto = f"        O objetivo dessa peça técnica é aferir os valores de mercado e de liquidação forçada por meio do método comparativo de dados de mercado, referente aos imóveis {imoveis}, localizados em #CIDADE_I - #ESTADO_I."

    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    runpar2 = par.add_run(texto)
    return doc

def texto_finalidade(doc):
    heading = doc.add_paragraph(style='Heading1')
    run = heading.add_run("3 - FINALIDADE")
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.font.name = "Cambria"
    run1 = doc.add_paragraph(" ")
    run2 = doc.add_paragraph("        Garantia bancária")

    return doc

def texto_proprietario(doc, lista_matriculas, lista_proponentes):
    heading = doc.add_paragraph( style='Heading1')
    run = heading.add_run("4 - PROPRIETÁRIO")
    run.font.name = "Cambria"
    run.font.color.rgb = RGBColor(0, 0, 0)
    run1 =doc.add_paragraph(" ")
    texto = gerar_texto_proprietarios_completo(lista_matriculas, lista_proponentes)
    run2 = doc.add_paragraph(texto)
    run2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return doc

def texto_ressalvas(doc):
    heading = doc.add_paragraph( style='Heading1')
    run = heading.add_run("5 - PRESSUPOSTOS, RESSALVAS E FATORES IMPORTANTES")
    run.font.name = "Cambria"
    run.font.color.rgb = RGBColor(0, 0, 0)
    run1 = doc.add_paragraph(" ")
    run2 = doc.add_paragraph("        Este Laudo fundamenta-se no que estabelecem as normas técnicas da ABNT"\
    "Avaliação de Bens, NBR 14653 – Parte 1 (Procedimentos Gerais/Revisão 2019) e Parte 3"\
    "(Imóveis Rurais/Revisão 2011), e baseia-se na documentação fornecida referente ao imóvel localizado"\
    "em #CIDADE_I - #ESTADO_I, situação na qual o #TRATAMENTO #SOLICITANTE solicita a avaliação do mesmo." \
    " Quanto às edificações e benfeitorias existentes no imóvel são considerados os quantitativos" \
    "de projetos existentes (se existirem), informações constatadas in loco quando da vistoria ao imóvel, " \
    "realizada em #DATA_VISTORIA e sendo, dessa forma, adotadas na presente avaliação como oficiais, por premissa," \
    " consideradas como válidas.")
    adicionar_espaco(doc)
    run3 = doc.add_paragraph("        Também, utilizamos como referência no decorrer dos trabalhos elementos documentais " \
    "e informações prestadas por terceiros, admitidas como confiáveis, corretas e de boa fé.")
    run2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    return doc