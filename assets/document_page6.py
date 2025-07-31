from docx import Document
from docx.shared import Pt, Cm
from docx.shared import RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Twips
import os
import win32com.client
from win32com.client.gencache import EnsureDispatch
from win32com.client import Dispatch
from assets.document_page5 import adicionar_espaco

def set_table_fixed_width(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    tblLayout = OxmlElement('w:tblLayout')
    tblLayout.set(qn('w:type'), 'fixed')
    tblPr.append(tblLayout)

def title_imovel(doc):
    heading = doc.add_paragraph( style='Heading 1')
    run = heading.add_run("6 - IDENTIFICAÇÃO E CARACTERIZAÇÃO DO IMÓVEL AVALIANDO")
    run.font.name = "Cambria"
    run.font.color.rgb = RGBColor(0, 0, 0)
    run1 = doc.add_paragraph("  \n  ")

    return doc

def localizacao(doc, lista_dados_matriculas):
    heading = doc.add_paragraph( style='Heading 2')
    run = heading.add_run("6.1 - Localização")
    run.font.name = "Cambria"
    run.font.color.rgb = RGBColor(0, 0, 0)

    locais_set = set()
    coords_dict = {}

    for dados in lista_dados_matriculas:
        cidade = dados.get("cidade", "").strip()
        estado = dados.get("estado", "").strip()
        latitude = dados.get("latitude", "").strip()
        longitude = dados.get("longitude", "").strip()
        matricula = dados.get("matricula", "").strip()
        nome_imovel = dados.get("nome_imovel", "").strip()

        if cidade and estado:
            locais_set.add(f"{cidade} - {estado}")

        if latitude and longitude:
            chave = (latitude, longitude)
            if chave not in coords_dict:
                coords_dict[chave] = []

            if (matricula, nome_imovel) not in coords_dict[chave]:
                coords_dict[chave].append((matricula, nome_imovel))

    print("🔎 Coordenadas agrupadas por centróide:")
    for (lat, lon), lista in coords_dict.items():
        print(f"  Matrícula(s): {[m for m, _ in lista]} | Imóvel(is): {[n for _, n in lista]} | Latitude: {lat} | Longitude: {lon}")

    texto_locais = ", ".join(sorted(locais_set))
    run1 = doc.add_paragraph(f"Zona Rural, Município: #CIDADE_I - #ESTADO_I")
    run2 = doc.add_paragraph("")

    for (lat, lon), lista in coords_dict.items():
        matriculas = [m for m, _ in lista]
        imoveis = [n for _, n in lista]
        if len(matriculas) == 1:
            par = doc.add_paragraph(f"Matrícula nº {matriculas[0]}, imóvel {imoveis[0]}")
            par.runs[0].underline = True
            par2 = doc.add_paragraph("Coordenadas:")
            par2.runs[0].underline = True
            par3 = doc.add_paragraph(f"Latitude: {lat}")
            par4 = doc.add_paragraph(f"Longitude: {lon}")
        else:
            matriculas_str = ", ".join(matriculas[:-1]) + " e " + matriculas[-1]
            imoveis_str = ", ".join(imoveis[:-1]) + " e " + imoveis[-1]
            par = doc.add_paragraph(f"Matrículas nº {matriculas_str}, imóveis {imoveis_str} respectivamente")
            par.runs[0].underline = True
            par2 = doc.add_paragraph("Coordenadas (centróide compartilhado):")
            par2.runs[0].underline = True
            par3 = doc.add_paragraph(f"Latitude: {lat}")
            par4 = doc.add_paragraph(f"Longitude: {lon}")

    return doc

def acesso(doc):
    heading = doc.add_paragraph( style='Heading 2')
    run = heading.add_run("6.2 - ROTA DE ACESSO ")
    run.font.name = "Cambria"
    run.font.color.rgb = RGBColor(0, 0, 0)

    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run1 = par.add_run("#IMAGEM_ACESSO")

    par2 = doc.add_paragraph()
    par2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run2 = par2.add_run("#ROTA_ACESSO")
    
    return doc

def carac_reg(doc):
    heading = doc.add_paragraph( style='Heading 2')
    run = heading.add_run("6.3 - Caracterização da Região")
    for run in heading.runs:
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0, 0, 0)
    run1 = doc.add_paragraph("#REGIAO_CIDADE \n #REGIAO_IMOVEL")
    run1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return doc

def desc_imovel(doc, lista_dados_matricula):
    heading = doc.add_paragraph( style='Heading 2')
    run = heading.add_run("6.4 - Descrição do imóvel")
    for run in heading.runs:
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0, 0, 0)
    
    for dados in lista_dados_matricula:
        matricula = dados.get("matricula", "")
        area_total = dados.get("area_total", "")
        p_reserva = dados.get("p_reserva", "")
        area_reserva = dados.get("area_reserva", '')
        p_app = dados.get("p_app", '')
        a_app = dados.get("a_app", '')
        observacoes = dados.get("observacoes_imovel",'')
        atividade_imovel = dados.get("atividade_imovel", "")

        run1 = doc.add_paragraph(f"Trata-se de um imóvel rural de Matrícula nº {matricula}, " \
        f"com área total de {area_total} ha, destes, {p_reserva} são separados para Reserva Legal, " \
        f"totalizando uma área de {area_reserva} ha, {observacoes}, a sua Área de Preservação Permanente – APP " \
        f"ocupa uma área {p_app}, totalizando {a_app} ha da integralidade do imóvel.")
        run1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

        run2 = doc.add_paragraph(f"{atividade_imovel}")
        run2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run3 = doc.add_paragraph("Uma melhor percepção do imóvel pode ser obtida através da tabela e das imagens a seguir:")
    run3.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    adicionar_espaco(doc)

    for i, dados in enumerate(lista_dados_matricula):
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        runpar = par.add_run(f"[INSERIR_TABELA_{i}_AQUI]")

        imagens_satelite(doc, lista_dados_matricula)

    adicionar_espaco(doc)
    return doc

def imagens_satelite(doc, lista_dados_matricula):
    for dados in lista_dados_matricula:
        matricula = dados.get("matricula", "")
        descricao_1 = dados.get("description_img_one", "Descrição da imagem 1")
        descricao_2 = dados.get("description_img_two", "Descrição da imagem 2")
        img_one = dados.get("img_one", "")
        img_two = dados.get("img_two", "")

        table2 = doc.add_table(rows=3,cols=2)
        table2.autofit
        table2.width = Cm(16)
        set_table_fixed_width(table2)
        table2.style = 'Table Grid'

        table2.rows[0].height = Cm(0.5)
        table2.rows[2].height = Cm(0.5)

        titlecell = table2.rows[0].cells[0]
        titlecell.merge(table2.rows[0].cells[1])
        titlecell.text = f"IMAGENS DE SATÉLITE DA PROPRIEDADE RURAL MATRICULA {matricula}"
        titlecell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        titlecell.paragraphs[0].runs[0].bold = True
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="4EA65D"/>')
        titlecell._tc.get_or_add_tcPr().append(shading)
        borders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>' 
            '<w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
            '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
            '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
            '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
            '</w:tcBorders>'
        )
        titlecell._tc.get_or_add_tcPr().append(borders)

        image_cell_1 = table2.rows[1].cells[0]
        p_img = image_cell_1.paragraphs[0]
        run_img = p_img.add_run()
        run_img.add_picture(f"{img_one}", width=Cm(8), height=Cm(4.84))
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER

        description_cell_1 = table2.rows[2].cells[0]
        description_cell_1.text = f"{descricao_1}"
        description_cell_1.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

        image_cell_2 = table2.rows[1].cells[1]
        p_img = image_cell_2.paragraphs[0]
        run_img = p_img.add_run()
        run_img.add_picture(f"{img_two}", width=Cm(8), height=Cm(4.84))
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER

        description_cell_2 = table2.rows[2].cells[1]
        description_cell_2.text = f"{descricao_2}"
        description_cell_2.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

def declividade(doc, imagem_path=None):
    heading = doc.add_paragraph( style='Heading 2')
    run = heading.add_run("6.4.1 - Declividade")
    for run in heading.runs:
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0, 0, 0)
    p = doc.add_paragraph()
    run1 = p.add_run("#IMAGEM_DECLIVIDADE")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    adicionar_espaco(doc)

    paragraph = doc.add_paragraph("a área apresenta declividade #DECLIVIDADE_I")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return doc

def hidrografia(doc, imagem_path=None):
    heading = doc.add_paragraph( style='Heading 2')
    run = heading.add_run("6.4.2 - Hidrografia")
    run.font.name = "Cambria"
    for run in heading.runs:
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0, 0, 0)
    p = doc.add_paragraph()
    run1 = p.add_run("#IMAGEM_HIDROGRAFIA")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    adicionar_espaco(doc)

    paragraph = doc.add_paragraph("a área apresenta hidrografia #HIDROGRAFIA_I")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return doc

def pedologia(doc, imagem_path=None):
    heading = doc.add_paragraph( style='Heading 2')
    run = heading.add_run("6.4.3 - Solo/Pedologia")
    run.font.name = "Cambria"
    for run in heading.runs:
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0, 0, 0)
    p = doc.add_paragraph()
    run1 = p.add_run("#IMAGEM_SOLOS")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    adicionar_espaco(doc)

    paragraph = doc.add_paragraph("a pedologia da região é predominada por #TIPO_SOLO")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    adicionar_espaco(doc)
    paragraph2 = doc.add_paragraph("a área apresenta #DESCRICAO_SOLO")
    paragraph2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return doc

def uso_imovel(doc, lista_dados_matriculas):
    heading = doc.add_paragraph( style='Heading 2')
    run = heading.add_run("6.4.4 – Potencial de Utilização do Imóvel")
    for run in heading.runs:
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0, 0, 0)
        run.font.name = "Cambria"
    for dados in lista_dados_matriculas:
        matricula = dados.get("matricula", "")
        potencial_imovel = dados.get("atividade_potencial", "")

        paragraph = doc.add_paragraph(f"O imóvel de matrícula {matricula} tem potencial de utilização para {potencial_imovel}.")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    return doc

def benfeitoria(doc, lista_dados_matriculas):
    heading = doc.add_paragraph( style='Heading 2')
    run = heading.add_run("6.4.5 – Benfeitorias")
    for run in heading.runs:
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0, 0, 0)
        run.font.name = "Cambria"

    for dados in lista_dados_matriculas:
        matricula = dados.get("matricula", "")
        nome_imovel = dados.get("nome_imovel", "")
        descricao_benfeitoria = dados.get("descricao_benfeitoria", "")
        area = dados.get("area", "")
        comprimento = dados.get("comprimento", "")
        largura = dados.get("largura", "")
        altura = dados.get("altura", "")
        imagem_benfeitoria = dados.get("imagem_benfeitoria", "")

        paragraph = doc.add_paragraph(f"No imóvel {nome_imovel} de matricula nº {matricula} foi observado a seguinte benfeitoria: ")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        table = doc.add_table(rows=4, cols=2)
        table.style = 'Table Grid'
        adicionar_espaco(doc)

        cell_1_1 = table.cell(0, 0)
        cell_1_2 = table.cell(0, 1)
        cell_1_1.text = 'BENFEITORIA'
        cell_1_1.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell_1_1.paragraphs[0].runs[0].bold = True
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="4EA65D"/>')
        cell_1_1._tc.get_or_add_tcPr().append(shading)
        borders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>' 
            '<w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
            '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
            '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
            '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
            '</w:tcBorders>'
        )
        cell_1_1._tc.get_or_add_tcPr().append(borders)

        cell_1_2.text = 'DESCRIÇÃO'
        cell_1_2.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell_1_2.paragraphs[0].runs[0].bold = True
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="4EA65D"/>')
        cell_1_2._tc.get_or_add_tcPr().append(shading)
        borders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            '<w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
            '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
            '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
            '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
            '</w:tcBorders>'
        )
        cell_1_2._tc.get_or_add_tcPr().append(borders)

        cell_benfeitoria = table.cell(1, 0).merge(table.cell(3, 0))

        cell_dimensoes = table.cell(1, 1)
        par = cell_dimensoes.paragraphs[0]
        par.add_run(f'Comprimento: {comprimento} m\n')
        par.add_run(f'Largura: {largura} m\n')
        par.add_run(f'Altura: {altura} m\n')
        par.add_run(f'Área Total: {area} m²')

        cell_desc = table.cell(2, 1)
        cell_desc.text = (
            f'{descricao_benfeitoria}.'
        )

        cell_ava = table.cell(3, 1)
        adicionar_estado_conservacao_com_tabela(cell_ava, lista_dados_matriculas)

        return doc

def adicionar_estado_conservacao_com_tabela(cell, lista_dados_matriculas):
    """Insere uma mini-tabela com estado de conservação dentro da célula"""
    for dados in lista_dados_matriculas:
        estado = dados.get("estado")

    p = cell.paragraphs[0]
    p.add_run('Estado de conservação:')

    nested_table = cell.add_table(rows=3, cols=7)
    nested_table.autofit = True
    nested_table.style = 'Table Grid'

    cell1_1 = nested_table.cell(0, 0)
    cell1_1.text = 'Bom'
    cell1_1.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    tcPr = cell1_1._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="nil"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="nil"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell1_2 = nested_table.cell(0, 1)
    cell1_2.text = ' '
    cell1_2.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    tcPr = cell1_2._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell1_3 = nested_table.cell(0, 2)
    cell1_3.text = ' '
    cell1_3.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    tcPr = cell1_3._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="nil"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell1_4 = nested_table.cell(0, 3)
    cell1_4.text = ' '
    cell1_4.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    tcPr = cell1_4._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="nil"/>'
        '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:right w:val="nil"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell1_5 = nested_table.cell(0, 4)
    cell1_5.text = 'Ruim'
    cell1_5.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    tcPr = cell1_5._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="nil"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell1_6 = nested_table.cell(0, 5)
    cell1_6.text = ' '
    cell1_6.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    tcPr = cell1_6._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell1_7 = nested_table.cell(0, 6)
    cell1_7.text = ' '
    cell1_7.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    tcPr = cell1_7._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="nil"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="nil"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell2_1 = nested_table.cell(1, 0)
    cell2_1.text = ' '
    cell2_1.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    tcPr = cell2_1._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="nil"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell2_2 = nested_table.cell(1, 1)
    cell2_2.text = ' '
    cell2_2.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    tcPr = cell2_2._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="nil"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="nil"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)
    cell2_3 = nested_table.cell(1, 2)
    cell2_3.text = 'Mediano'
    cell2_3.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER 
    tcPr = cell2_3._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="nil"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="nil"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell2_4 = nested_table.cell(1, 3)
    cell2_4.text = ' '
    cell2_4.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    tcPr = cell2_4._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell2_5 = nested_table.cell(1, 4)
    cell2_5.text = ' '
    cell2_5.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    tcPr = cell2_5._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="nil"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="nil"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell2_6 = nested_table.cell(1, 5)
    cell2_6.text = ' '
    cell2_6.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    tcPr = cell2_6._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="nil"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="nil"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell2_7 = nested_table.cell(1, 6)
    cell2_7.text = ' '
    cell2_7.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    tcPr = cell2_7._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="nil"/>'
        '<w:bottom w:val="nil"/>'
        '<w:right w:val="nil"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    cell3_1 = nested_table.cell(2, 0)
    cell3_1.merge(nested_table.cell(2, 6))
    cell3_1.text = ' '
    cell3_1.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    tcPr = cell3_1._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        '<w:top w:val="nil"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

    if estado == "bom":
        target_cell = cell1_2
    elif estado == "mediano":
        target_cell = cell2_4
    elif estado == "ruim":
        target_cell = cell1_6
    else:
        target_cell = None

    if target_cell:
        shade = parse_xml(r'<w:shd {} w:fill="4EA65D"/>'.format(nsdecls('w')))
        target_cell._tc.get_or_add_tcPr().append(shade)
