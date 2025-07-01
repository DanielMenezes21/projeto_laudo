from docx import Document
from docx.shared import Cm

def criar_tabela_benfeitoria(doc: Document):
    # Adiciona um título
    doc.add_heading('BENFEITORIA', level=1)

    # Adiciona uma tabela com 2 colunas
    table = doc.add_table(rows=0, cols=2)
    table.style = 'Table Grid'

    # Descrição
    row = table.add_row().cells
    row[0].text = 'DESCRIÇÃO'
    row[1].text = ''

    # Dimensões
    row = table.add_row().cells
    row[0].text = 'Comprimento'
    row[1].text = '20,0 m'

    row = table.add_row().cells
    row[0].text = 'Largura'
    row[1].text = '10,0 m'

    row = table.add_row().cells
    row[0].text = 'Altura'
    row[1].text = '10,0 m'

    row = table.add_row().cells
    row[0].text = 'Área Total'
    row[1].text = '200,0 m²'

    # Finalidade
    doc.add_paragraph(
        'O galpão é ideal para armazenamento e organização de bags de fertilizantes e sementes '
        'para plantio, além de alocar os maquinários da propriedade.'
    )

    # Estado de Conservação
    doc.add_heading('Estado de Conservação:', level=2)

    # Marcação de opções
    estados = ['Bom', 'Mediano', 'Ruim']
    for estado in estados:
        doc.add_paragraph(estado, style='List Bullet')

    return doc

# Exemplo de uso
doc = Document()
criar_tabela_benfeitoria(doc)
doc.save('tabela_gerada.docx')
