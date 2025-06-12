from docx import Document
from docx.shared import Pt, Cm
from docx.shared import RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Twips

def adicionar_espaco(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)  
    p.paragraph_format.space_after = Pt(0)   
    run = p.add_run(" ")
    run.font.name = 'Calibri'
    run.font.size = Pt(12) 

    return doc

def texto_solicitante(doc):
    heading = doc.add_heading("1 - SOLICITANTE", level=1)
    run = heading.runs[0]
    run.font.color.rgb = RGBColor(0, 0, 0)
    doc.add_paragraph("Fomos solicitados pelo {prop} {nome}, para avaliar um imóvel rural, denominado {imovel}, localizado em {cid_est} ")

    return doc
