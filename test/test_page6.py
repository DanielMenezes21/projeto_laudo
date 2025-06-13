from docx import Document
from docx.shared import Pt, Cm
from docx.shared import RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Twips

def title_imovel(doc):
    heading = doc.add_heading("6 - IDENTIFICAÇÃO E CARACTERIZAÇÃO DO IMÓVEL AVALIANDO", level=1)
    run = heading.runs[0]
    run.font.color.rgb = RGBColor(0, 0, 0)
    doc.add_paragraph("  \n  ")

    return doc

def localizacao(doc):
    heading = doc.add_heading("6.1 - Localização", level=2)
    run = heading.runs[0]
    run.font.color.rgb = RGBColor(0, 0, 0)
    run1 = doc.add_paragraph("Zona Rural, Município de {cid_est}")
    run2 = doc.add_paragraph("")
    run3 = doc.add_paragraph("Coordenadas: ")
    run3.underline = True
    run4 = doc.add_paragraph("Latitude: {latitude}")
    run5 = doc.add_paragraph("Longitude: {Longitude}")
    run2
    run2
    for i in [run1, run2, run3, run4, run5]:
        for r in i.runs:
            r.font.size = Pt(12)
    return doc

def acesso(doc, imagem_acesso=None):
    heading = doc.add_heading("6.2 - ROTA DE ACESSO MATRÌCULA {mat}", level=2)
    run = heading.runs[0]
    for run in heading.runs:
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0, 0, 0)
    run1 = doc.add_paragraph("{rota_acesso}")
    for r in run1.runs:
        r.font.size = Pt(12)
    if imagem_acesso:
        try:
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p_img.add_run()
            shape = run.add_picture(imagem_acesso, width=Cm(15), height=Cm(8.77))
            border_xml = (
                '<pic:pic xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture" '
                'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
                '<pic:spPr>'
                '<a:ln w="9525">'  
                '<a:solidFill>'
                '<a:srgbClr val="000000"/>'  
                '</a:solidFill>'
                '<a:prstDash val="solid"/>'  
                '</a:ln>'
                '</pic:spPr>'
                '</pic:pic>'
            )

            pic = run._r.xpath('.//pic:pic')[0]
            pic.append(parse_xml(border_xml))
            
        except Exception as e:
            print(f"Erro ao adicionar imagem: {e}")
            doc.add_paragraph("[ESPAÇO PARA ACESSO]", style='Normal')
    else:
        doc.add_paragraph("[ESPAÇO PARA ACESSO]", style='Normal')

    return doc

def carac_reg(doc):
    heading = doc.add_heading("6.3 - Caracterização da Região", level=2)
    run = heading.runs[0]
    for run in heading.runs:
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0, 0, 0)
    run1 = doc.add_paragraph("{caracterização}")
    return doc

def desc_imovel(doc):
    heading = doc.add_heading("6.4 - Descrição do imóvel", level=2)
    run = heading.runs[0]
    for run in heading.runs:
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0, 0, 0)
    run1 = doc.add_paragraph("Trata-se de um imóvel rural de Matrícula nº 154.725, " \
    "com área total de {area_total} ha, destes, {p_reserva} são separados para Reserva Legal, " \
    "totalizando uma área de {area_reserva} ha, {observacao}, a sua Área de Preservação Permanente – APP " \
    "ocupa uma área {p_app}, totalizando {area_app} ha da integralidade do imóvel.")
    run2 = doc.add_paragraph("{descricao_atividade}")
    run3 = doc.add_paragraph("Uma melhor percepção do imóvel pode ser obtida através da tabela e das imagens a seguir:")

    table = doc.add_table(row=1,col=1)
    table.allow_autofit = False
    table.width = Cm(17)
    table.style = 'Table Grid'

    table.rows[0].height = Cm(0.5)

    



